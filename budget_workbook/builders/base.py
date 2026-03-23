"""Base classes and helper utilities for worksheet builders."""

from __future__ import annotations

from abc import ABC, abstractmethod
from collections.abc import Callable

from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from budget_workbook.config import WorkbookConfig
from budget_workbook.rows import MonthlyEntryRows
from budget_workbook.styles import WorkbookStyles


FormulaFactory = Callable[[int, str], str]


class BaseSheetBuilder(ABC):
    """Base class for every worksheet builder."""

    label_col = 2
    month_start_col = 3
    group_width = 3

    def __init__(
        self,
        config: WorkbookConfig,
        styles: WorkbookStyles,
        rows: MonthlyEntryRows,
    ) -> None:
        self.config = config
        self.styles = styles
        self.rows = rows

    @abstractmethod
    def build(self, workbook):
        """Create and return the worksheet."""

    def month_budget_col(self, index: int) -> int:
        return self.month_start_col + index * self.group_width

    def month_actual_col(self, index: int) -> int:
        return self.month_budget_col(index) + 1

    def month_variance_col(self, index: int) -> int:
        return self.month_budget_col(index) + 2

    def year_budget_col(self) -> int:
        return self.month_start_col + len(self.config.months) * self.group_width

    def year_actual_col(self) -> int:
        return self.year_budget_col() + 1

    def year_variance_col(self) -> int:
        return self.year_budget_col() + 2

    def last_data_col(self) -> int:
        return self.year_variance_col()

    def month_budget_letter(self, index: int) -> str:
        return get_column_letter(self.month_budget_col(index))

    def month_actual_letter(self, index: int) -> str:
        return get_column_letter(self.month_actual_col(index))

    def month_variance_letter(self, index: int) -> str:
        return get_column_letter(self.month_variance_col(index))

    def year_budget_letter(self) -> str:
        return get_column_letter(self.year_budget_col())

    def year_actual_letter(self) -> str:
        return get_column_letter(self.year_actual_col())

    def year_variance_letter(self) -> str:
        return get_column_letter(self.year_variance_col())

    def apply_border_range(
        self,
        worksheet: Worksheet,
        start_row: int,
        end_row: int,
        start_col: int,
        end_col: int,
    ) -> None:
        for row in range(start_row, end_row + 1):
            for col in range(start_col, end_col + 1):
                worksheet.cell(row=row, column=col).border = self.styles.thin_border

    def create_section_header(
        self,
        worksheet: Worksheet,
        row: int,
        title: str,
        start_col: int | None = None,
        end_col: int | None = None,
    ) -> None:
        first_col = start_col or self.label_col
        last_col = end_col or self.last_data_col()
        worksheet.merge_cells(
            start_row=row,
            start_column=first_col,
            end_row=row,
            end_column=last_col,
        )
        cell = worksheet.cell(row=row, column=first_col)
        cell.value = title
        cell.font = self.styles.section_font()
        cell.fill = self.styles.section_fill()
        cell.alignment = self.styles.left_alignment()
        self.apply_border_range(worksheet, row, row, first_col, last_col)
        worksheet.row_dimensions[row].height = 22

    def create_month_header_grid(self, worksheet: Worksheet, title_row: int, subtitle_row: int) -> None:
        worksheet[f"B{title_row}"] = "CATEGORY"
        worksheet[f"B{title_row}"].font = self.styles.header_font()
        worksheet[f"B{title_row}"].fill = self.styles.header_fill()
        worksheet[f"B{title_row}"].border = self.styles.thin_border
        worksheet[f"B{title_row}"].alignment = self.styles.left_alignment()
        worksheet.merge_cells(start_row=title_row, start_column=2, end_row=subtitle_row, end_column=2)

        for index, month in enumerate(self.config.months):
            start_col = self.month_budget_col(index)
            worksheet.merge_cells(
                start_row=title_row,
                start_column=start_col,
                end_row=title_row,
                end_column=start_col + 2,
            )
            month_cell = worksheet.cell(row=title_row, column=start_col)
            month_cell.value = month
            month_cell.font = self.styles.header_font()
            month_cell.fill = self.styles.header_fill()
            month_cell.alignment = self.styles.center_alignment()
            self.apply_border_range(worksheet, title_row, title_row, start_col, start_col + 2)

            for offset, label in enumerate(("Budget", "Actual", "Variance")):
                cell = worksheet.cell(row=subtitle_row, column=start_col + offset)
                cell.value = label
                cell.font = self.styles.header_font()
                cell.fill = self.styles.header_fill()
                cell.border = self.styles.thin_border
                cell.alignment = self.styles.center_alignment()

        start_col = self.year_budget_col()
        worksheet.merge_cells(
            start_row=title_row,
            start_column=start_col,
            end_row=title_row,
            end_column=start_col + 2,
        )
        total_cell = worksheet.cell(row=title_row, column=start_col)
        total_cell.value = "YEAR TOTAL"
        total_cell.font = self.styles.header_font()
        total_cell.fill = self.styles.header_fill()
        total_cell.alignment = self.styles.center_alignment()
        self.apply_border_range(worksheet, title_row, title_row, start_col, start_col + 2)

        for offset, label in enumerate(("Budget", "Actual", "Variance")):
            cell = worksheet.cell(row=subtitle_row, column=start_col + offset)
            cell.value = label
            cell.font = self.styles.header_font()
            cell.fill = self.styles.header_fill()
            cell.border = self.styles.thin_border
            cell.alignment = self.styles.center_alignment()

        worksheet.row_dimensions[title_row].height = 25
        worksheet.row_dimensions[subtitle_row].height = 22

    def create_input_row(
        self,
        worksheet: Worksheet,
        row: int,
        label: str,
        *,
        indent: bool = False,
        use_percent: bool = False,
    ) -> None:
        self._set_row_label(worksheet, row, label, indent=indent)
        for index, _month in enumerate(self.config.months):
            budget_cell = worksheet.cell(row=row, column=self.month_budget_col(index))
            actual_cell = worksheet.cell(row=row, column=self.month_actual_col(index))
            variance_cell = worksheet.cell(row=row, column=self.month_variance_col(index))
            for cell in (budget_cell, actual_cell, variance_cell):
                cell.border = self.styles.thin_border
                cell.alignment = self.styles.right_alignment()
                cell.number_format = self.config.percent_format if use_percent else self.config.currency_format
            budget_cell.font = self.styles.input_font()
            actual_cell.font = self.styles.input_font()
            variance_cell.font = self.styles.calc_font()
            variance_cell.value = (
                f"=IF(OR({get_column_letter(budget_cell.column)}{row}<>\"\","
                f"{get_column_letter(actual_cell.column)}{row}<>\"\"),"
                f"{get_column_letter(actual_cell.column)}{row}-{get_column_letter(budget_cell.column)}{row},0)"
            )

        self._set_year_total_formulas(worksheet, row, use_percent=use_percent)
        self.apply_alternating_fill(worksheet, row)

    def create_formula_row(
        self,
        worksheet: Worksheet,
        row: int,
        label: str,
        budget_formula_factory: FormulaFactory,
        actual_formula_factory: FormulaFactory,
        *,
        indent: bool = False,
        use_percent: bool = False,
        fill=None,
        font_color: str | None = None,
        bold: bool = False,
    ) -> None:
        self._set_row_label(
            worksheet,
            row,
            label,
            indent=indent,
            fill=fill,
            font_color=font_color,
            bold=bold,
        )
        for index, month in enumerate(self.config.months):
            budget_col = self.month_budget_col(index)
            actual_col = self.month_actual_col(index)
            variance_col = self.month_variance_col(index)

            budget_cell = worksheet.cell(row=row, column=budget_col)
            actual_cell = worksheet.cell(row=row, column=actual_col)
            variance_cell = worksheet.cell(row=row, column=variance_col)

            budget_cell.value = budget_formula_factory(index, month)
            actual_cell.value = actual_formula_factory(index, month)
            variance_cell.value = f"={get_column_letter(actual_col)}{row}-{get_column_letter(budget_col)}{row}"

            for cell in (budget_cell, actual_cell, variance_cell):
                cell.border = self.styles.thin_border
                cell.alignment = self.styles.right_alignment()
                cell.font = Font(
                    color=font_color or self.styles.palette.text_dark,
                    bold=bold,
                    size=11,
                )
                cell.number_format = self.config.percent_format if use_percent else self.config.currency_format
                if fill is not None:
                    cell.fill = fill

        self._set_year_total_formulas(
            worksheet,
            row,
            use_percent=use_percent,
            fill=fill,
            font_color=font_color,
            bold=bold,
        )
        if fill is None:
            self.apply_alternating_fill(worksheet, row)

    def apply_alternating_fill(self, worksheet: Worksheet, row: int) -> None:
        if row % 2 == 0:
            for col in range(self.label_col, self.last_data_col() + 1):
                worksheet.cell(row=row, column=col).fill = self.styles.alt_row_fill()

    def add_variance_conditional_formatting(
        self,
        worksheet: Worksheet,
        start_row: int,
        end_row: int,
        columns: list[int] | None = None,
    ) -> None:
        from openpyxl.formatting.rule import CellIsRule
        from openpyxl.styles import PatternFill

        red_fill = PatternFill(
            start_color=self.styles.palette.red_fill,
            end_color=self.styles.palette.red_fill,
            fill_type="solid",
        )
        green_fill = PatternFill(
            start_color=self.styles.palette.green_fill,
            end_color=self.styles.palette.green_fill,
            fill_type="solid",
        )
        variance_columns = columns or [self.month_variance_col(index) for index in range(len(self.config.months))]
        if columns is None:
            variance_columns.append(self.year_variance_col())
        for col in variance_columns:
            cell_range = f"{get_column_letter(col)}{start_row}:{get_column_letter(col)}{end_row}"
            worksheet.conditional_formatting.add(
                cell_range,
                CellIsRule(operator="greaterThan", formula=["0"], fill=red_fill),
            )
            worksheet.conditional_formatting.add(
                cell_range,
                CellIsRule(operator="lessThan", formula=["0"], fill=green_fill),
            )

    def _set_row_label(
        self,
        worksheet: Worksheet,
        row: int,
        label: str,
        *,
        indent: bool = False,
        fill=None,
        font_color: str | None = None,
        bold: bool = False,
    ) -> None:
        label_cell = worksheet.cell(row=row, column=self.label_col)
        label_cell.value = f"  {label}" if indent else label
        label_cell.font = Font(
            size=10 if not bold else 11,
            bold=bold or not indent,
            color=font_color or self.styles.palette.text_dark,
        )
        label_cell.alignment = self.styles.left_alignment()
        label_cell.border = self.styles.thin_border
        if fill is not None:
            label_cell.fill = fill

    def _set_year_total_formulas(
        self,
        worksheet: Worksheet,
        row: int,
        *,
        use_percent: bool,
        fill=None,
        font_color: str | None = None,
        bold: bool = False,
    ) -> None:
        budget_sum_parts = []
        actual_sum_parts = []
        for index in range(len(self.config.months)):
            budget_sum_parts.append(f"{self.month_budget_letter(index)}{row}")
            actual_sum_parts.append(f"{self.month_actual_letter(index)}{row}")

        year_budget = worksheet.cell(row=row, column=self.year_budget_col())
        year_actual = worksheet.cell(row=row, column=self.year_actual_col())
        year_variance = worksheet.cell(row=row, column=self.year_variance_col())

        year_budget.value = f"=SUM({','.join(budget_sum_parts)})"
        year_actual.value = f"=SUM({','.join(actual_sum_parts)})"
        year_variance.value = f"={self.year_actual_letter()}{row}-{self.year_budget_letter()}{row}"

        for cell in (year_budget, year_actual, year_variance):
            cell.border = self.styles.thin_border
            cell.alignment = self.styles.right_alignment()
            cell.font = Font(
                color=font_color or self.styles.palette.text_dark,
                bold=bold,
                size=11,
            )
            cell.number_format = self.config.percent_format if use_percent else self.config.currency_format
            if fill is not None:
                cell.fill = fill
