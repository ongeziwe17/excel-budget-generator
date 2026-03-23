"""Bonus tracker worksheet builder."""

from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from .base import BaseSheetBuilder


class BonusTrackerSheetBuilder(BaseSheetBuilder):
    """Build the Bonus Tracker worksheet."""

    def build(self, workbook):
        worksheet = workbook.create_sheet("Bonus Tracker")
        worksheet.sheet_view.showGridLines = False

        worksheet.column_dimensions["A"].width = 3
        worksheet.column_dimensions["B"].width = 25
        worksheet.column_dimensions["C"].width = 18
        worksheet.column_dimensions["D"].width = 18
        worksheet.column_dimensions["E"].width = 25

        worksheet.merge_cells("B2:E2")
        worksheet["B2"] = "BONUS TRACKER"
        worksheet["B2"].font = Font(size=18, bold=True, color=self.styles.palette.header_dark)
        worksheet["B2"].alignment = self.styles.left_alignment()
        worksheet.row_dimensions[2].height = 30

        worksheet.merge_cells("B3:E3")
        worksheet["B3"] = "Track your quarterly bonuses - enter actual amounts received"
        worksheet["B3"].font = Font(size=10, italic=True, color="666666")

        worksheet["B5"] = "EXPECTED BONUS SCHEDULE (based on your profile)"
        worksheet["B5"].font = Font(size=12, bold=True, color=self.styles.palette.header_dark)

        for cell_ref, label, alignment in [
            ("B6", "Quarter", self.styles.left_alignment()),
            ("C6", "% of Gross", self.styles.center_alignment()),
            ("D6", "Expected Amount", self.styles.right_alignment()),
            ("E6", "Actual Received", self.styles.right_alignment()),
        ]:
            worksheet[cell_ref] = label
            worksheet[cell_ref].font = self.styles.header_font()
            worksheet[cell_ref].fill = self.styles.header_fill()
            worksheet[cell_ref].border = self.styles.thin_border
            worksheet[cell_ref].alignment = alignment
        worksheet.row_dimensions[6].height = 22

        bonus_schedule = [
            ("Q1 (Jan-Mar)", "0%", 0),
            ("Q2 (Apr-Jun)", "25%", self.config.monthly_gross_salary * self.config.bonus_q2_pct),
            ("Q3 (Jul-Sep)", "50%", self.config.monthly_gross_salary * self.config.bonus_q3_pct),
            ("Q4 (Oct-Dec)", "83%", self.config.monthly_gross_salary * self.config.bonus_q4_pct),
        ]

        row = 7
        for quarter, percent_text, expected in bonus_schedule:
            worksheet[f"B{row}"] = quarter
            worksheet[f"B{row}"].font = Font(size=10, color=self.styles.palette.text_dark)
            worksheet[f"B{row}"].border = self.styles.thin_border

            worksheet[f"C{row}"] = percent_text
            worksheet[f"C{row}"].font = Font(size=10, color=self.styles.palette.text_dark)
            worksheet[f"C{row}"].border = self.styles.thin_border
            worksheet[f"C{row}"].alignment = self.styles.center_alignment()

            worksheet[f"D{row}"] = expected
            worksheet[f"D{row}"].font = Font(size=10, color="666666")
            worksheet[f"D{row}"].border = self.styles.thin_border
            worksheet[f"D{row}"].alignment = self.styles.right_alignment()
            worksheet[f"D{row}"].number_format = self.config.currency_format

            worksheet[f"E{row}"].border = self.styles.thin_border
            worksheet[f"E{row}"].alignment = self.styles.right_alignment()
            worksheet[f"E{row}"].number_format = self.config.currency_format
            worksheet[f"E{row}"].font = self.styles.input_font()

            if row % 2 == 0:
                for col in ["B", "C", "D", "E"]:
                    worksheet[f"{col}{row}"].fill = self.styles.alt_row_fill()
            row += 1

        worksheet[f"B{row}"] = "TOTAL YEARLY BONUS"
        worksheet[f"B{row}"].font = Font(bold=True, color=self.styles.palette.white)
        worksheet[f"B{row}"].fill = self.styles.header_fill()
        worksheet[f"B{row}"].border = self.styles.thin_border

        worksheet[f"C{row}"].border = self.styles.thin_border
        worksheet[f"C{row}"].fill = self.styles.header_fill()

        total_expected = self.config.monthly_gross_salary * (
            self.config.bonus_q2_pct + self.config.bonus_q3_pct + self.config.bonus_q4_pct
        )
        worksheet[f"D{row}"] = total_expected
        worksheet[f"D{row}"].font = Font(bold=True, color=self.styles.palette.white)
        worksheet[f"D{row}"].fill = self.styles.header_fill()
        worksheet[f"D{row}"].border = self.styles.thin_border
        worksheet[f"D{row}"].alignment = self.styles.right_alignment()
        worksheet[f"D{row}"].number_format = self.config.currency_format

        worksheet[f"E{row}"] = "=SUM(E7:E10)"
        worksheet[f"E{row}"].font = Font(bold=True, color=self.styles.palette.white)
        worksheet[f"E{row}"].fill = self.styles.header_fill()
        worksheet[f"E{row}"].border = self.styles.thin_border
        worksheet[f"E{row}"].alignment = self.styles.right_alignment()
        worksheet[f"E{row}"].number_format = self.config.currency_format

        row += 3
        worksheet[f"B{row}"] = "VARIANCE FROM EXPECTED"
        worksheet[f"B{row}"].font = Font(size=12, bold=True, color=self.styles.palette.header_dark)

        row += 2
        worksheet[f"B{row}"] = "Difference (Actual - Expected)"
        worksheet[f"B{row}"].font = Font(size=10, color=self.styles.palette.text_dark)
        worksheet[f"B{row}"].border = self.styles.thin_border

        worksheet[f"E{row}"] = "=E11-D11"
        worksheet[f"E{row}"].font = self.styles.calc_font()
        worksheet[f"E{row}"].border = self.styles.thin_border
        worksheet[f"E{row}"].alignment = self.styles.right_alignment()
        worksheet[f"E{row}"].number_format = self.config.currency_format

        row += 3
        worksheet[f"B{row}"] = "MONTHLY BONUS TRACKING (from Monthly Entry)"
        worksheet[f"B{row}"].font = Font(size=12, bold=True, color=self.styles.palette.header_dark)

        row += 2
        worksheet[f"B{row}"] = "Month"
        worksheet[f"C{row}"] = "Bonus Amount"
        worksheet[f"B{row}"].font = self.styles.header_font()
        worksheet[f"C{row}"].font = self.styles.header_font()
        worksheet[f"B{row}"].fill = self.styles.header_fill()
        worksheet[f"C{row}"].fill = self.styles.header_fill()
        worksheet[f"B{row}"].border = self.styles.thin_border
        worksheet[f"C{row}"].border = self.styles.thin_border
        worksheet[f"C{row}"].alignment = self.styles.right_alignment()
        worksheet.row_dimensions[row].height = 22
        row += 1

        for index, month in enumerate(self.config.months):
            col = get_column_letter(3 + index)
            worksheet[f"B{row}"] = month
            worksheet[f"B{row}"].font = Font(size=10, color=self.styles.palette.text_dark)
            worksheet[f"B{row}"].border = self.styles.thin_border

            worksheet[f"C{row}"] = f"='Monthly Entry'!{col}{self.rows.bonus}"
            worksheet[f"C{row}"].font = self.styles.calc_font()
            worksheet[f"C{row}"].border = self.styles.thin_border
            worksheet[f"C{row}"].alignment = self.styles.right_alignment()
            worksheet[f"C{row}"].number_format = self.config.currency_format

            if row % 2 == 0:
                worksheet[f"B{row}"].fill = self.styles.alt_row_fill()
                worksheet[f"C{row}"].fill = self.styles.alt_row_fill()
            row += 1

        worksheet[f"B{row}"] = "Total"
        worksheet[f"B{row}"].font = Font(bold=True, color=self.styles.palette.text_dark)
        worksheet[f"B{row}"].border = self.styles.thin_border

        worksheet[f"C{row}"] = "='Monthly Entry'!O36"
        worksheet[f"C{row}"].font = Font(bold=True, color=self.styles.palette.text_dark)
        worksheet[f"C{row}"].border = self.styles.thin_border
        worksheet[f"C{row}"].alignment = self.styles.right_alignment()
        worksheet[f"C{row}"].number_format = self.config.currency_format

        return worksheet
