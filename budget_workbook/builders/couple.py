"""Worksheet builders for the two-person household budget mode."""

from __future__ import annotations

from dataclasses import dataclass

from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.chart.series import SeriesLabel
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo

from budget_workbook.config import ContributionMethod, PersonConfig

from .base import BaseSheetBuilder


@dataclass(frozen=True)
class PartnerBudgetRows:
    """Stable rows shared by both partner budget sheets."""

    net_income: int = 8
    other_income: int = 9
    bonus: int = 10
    fixed_expenses: int = 13
    variable_expenses: int = 14
    total_expenses: int = 15
    personal_savings: int = 18
    household_contribution: int = 19
    opening_savings: int = 20
    total_income: int = 23
    disposable_balance: int = 24
    savings_rate: int = 25
    closing_savings: int = 26


@dataclass(frozen=True)
class SharedBudgetRows:
    """Stable rows in the shared household budget sheet."""

    partner_one_contribution: int = 8
    partner_two_contribution: int = 9
    total_contributions: int = 10
    fixed_expenses: int = 13
    variable_expenses: int = 14
    total_expenses: int = 15
    joint_savings: int = 18
    opening_joint_savings: int = 19
    required_funding: int = 22
    shared_balance: int = 23
    closing_joint_savings: int = 24


PARTNER_ROWS = PartnerBudgetRows()
SHARED_ROWS = SharedBudgetRows()


class CoupleBaseSheetBuilder(BaseSheetBuilder):
    """Common helpers for household-mode worksheets."""

    def configure_monthly_columns(self, worksheet, label_width: int = 34) -> None:
        worksheet.column_dimensions["A"].width = 3
        worksheet.column_dimensions["B"].width = label_width
        for index in range(len(self.config.months)):
            for col in (
                self.month_budget_col(index),
                self.month_actual_col(index),
                self.month_variance_col(index),
            ):
                worksheet.column_dimensions[get_column_letter(col)].width = 12
        for col in (self.year_budget_col(), self.year_actual_col(), self.year_variance_col()):
            worksheet.column_dimensions[get_column_letter(col)].width = 14

    def title(self, worksheet, text: str, subtitle: str, end_col: int | None = None) -> None:
        final_col = end_col or self.last_data_col()
        worksheet.merge_cells(start_row=2, start_column=2, end_row=2, end_column=final_col)
        worksheet["B2"] = text
        worksheet["B2"].font = Font(size=18, bold=True, color=self.styles.palette.header_dark)
        worksheet["B2"].alignment = self.styles.left_alignment()
        worksheet.row_dimensions[2].height = 30
        worksheet.merge_cells(start_row=3, start_column=2, end_row=3, end_column=final_col)
        worksheet["B3"] = subtitle
        worksheet["B3"].font = Font(size=10, italic=True, color="666666")

    def link_row(
        self,
        worksheet,
        row: int,
        label: str,
        budget_factory,
        actual_factory,
        *,
        percent: bool = False,
        total_last_value: bool = False,
    ) -> None:
        self._set_row_label(worksheet, row, label)
        for index, month in enumerate(self.config.months):
            budget_col = self.month_budget_col(index)
            actual_col = self.month_actual_col(index)
            variance_col = self.month_variance_col(index)
            worksheet.cell(row=row, column=budget_col).value = budget_factory(index, month)
            worksheet.cell(row=row, column=actual_col).value = actual_factory(index, month)
            worksheet.cell(row=row, column=variance_col).value = (
                f"={get_column_letter(actual_col)}{row}-{get_column_letter(budget_col)}{row}"
            )
            for col in (budget_col, actual_col, variance_col):
                cell = worksheet.cell(row=row, column=col)
                cell.border = self.styles.thin_border
                cell.alignment = self.styles.right_alignment()
                cell.number_format = self.config.percent_format if percent else self.config.currency_format
                cell.font = self.styles.calc_font()

        if total_last_value:
            last_index = len(self.config.months) - 1
            worksheet.cell(row=row, column=self.year_budget_col()).value = (
                f"={self.month_budget_letter(last_index)}{row}"
            )
            worksheet.cell(row=row, column=self.year_actual_col()).value = (
                f"={self.month_actual_letter(last_index)}{row}"
            )
        elif percent:
            worksheet.cell(row=row, column=self.year_budget_col()).value = (
                f"=IF({self.year_budget_letter()}{PARTNER_ROWS.total_income}=0,0,"
                f"{self.year_budget_letter()}{PARTNER_ROWS.personal_savings}/"
                f"{self.year_budget_letter()}{PARTNER_ROWS.total_income})"
            )
            worksheet.cell(row=row, column=self.year_actual_col()).value = (
                f"=IF({self.year_actual_letter()}{PARTNER_ROWS.total_income}=0,0,"
                f"{self.year_actual_letter()}{PARTNER_ROWS.personal_savings}/"
                f"{self.year_actual_letter()}{PARTNER_ROWS.total_income})"
            )
        else:
            budget_cells = [f"{self.month_budget_letter(i)}{row}" for i in range(len(self.config.months))]
            actual_cells = [f"{self.month_actual_letter(i)}{row}" for i in range(len(self.config.months))]
            worksheet.cell(row=row, column=self.year_budget_col()).value = f"=SUM({','.join(budget_cells)})"
            worksheet.cell(row=row, column=self.year_actual_col()).value = f"=SUM({','.join(actual_cells)})"
        worksheet.cell(row=row, column=self.year_variance_col()).value = (
            f"={self.year_actual_letter()}{row}-{self.year_budget_letter()}{row}"
        )
        for col in (self.year_budget_col(), self.year_actual_col(), self.year_variance_col()):
            cell = worksheet.cell(row=row, column=col)
            cell.border = self.styles.thin_border
            cell.alignment = self.styles.right_alignment()
            cell.number_format = self.config.percent_format if percent else self.config.currency_format
            cell.font = self.styles.calc_font()

    def add_input_variance_rules(self, worksheet, start_row: int, end_row: int) -> None:
        self.add_variance_conditional_formatting(worksheet, start_row, end_row)


class HouseholdCoverSheetBuilder(CoupleBaseSheetBuilder):
    """Build the household workbook cover and usage guidance."""

    def build(self, workbook):
        worksheet = workbook.active
        worksheet.title = "Cover"
        worksheet.sheet_view.showGridLines = False
        for col, width in zip(["A", "B", "C", "D", "E", "F"], [3, 28, 24, 24, 24, 3]):
            worksheet.column_dimensions[col].width = width
        worksheet.merge_cells("B3:E3")
        worksheet["B3"] = "COUPLE & HOUSEHOLD BUDGET"
        worksheet["B3"].font = Font(size=20, bold=True, color=self.styles.palette.header_dark)
        worksheet["B3"].alignment = self.styles.center_alignment()
        worksheet.merge_cells("B5:E5")
        worksheet["B5"] = "='Household Setup'!C7&\" and \"&'Household Setup'!D7"
        worksheet["B5"].font = Font(size=14, bold=True, color=self.styles.palette.header_light)
        worksheet["B5"].alignment = self.styles.center_alignment()
        worksheet.merge_cells("B8:E8")
        worksheet["B8"] = "How money flows through this workbook"
        worksheet["B8"].font = Font(size=14, bold=True, color=self.styles.palette.header_dark)
        instructions = [
            "1. Review partner names, income assumptions and the shared-cost split on Household Setup.",
            "2. Enter personal actuals on each Partner Budget sheet.",
            "3. Enter shared household expenses and joint savings on Shared Household.",
            "4. Review combined results on Household Dashboard, then use the supporting trackers.",
            "Household contributions are internal transfers and are excluded from consolidated income and expenses.",
        ]
        for row, instruction in enumerate(instructions, start=10):
            worksheet.merge_cells(start_row=row, start_column=2, end_row=row, end_column=5)
            worksheet.cell(row=row, column=2).value = instruction
            worksheet.cell(row=row, column=2).font = Font(size=11, color=self.styles.palette.text_dark)
        worksheet["B18"] = "Blue text"
        worksheet["B18"].font = self.styles.input_font()
        worksheet["C18"] = "Editable workbook inputs"
        worksheet["B19"] = "Black text"
        worksheet["C19"] = "Calculated values"
        return worksheet


class HouseholdSetupSheetBuilder(CoupleBaseSheetBuilder):
    """Build household assumptions and contribution allocation controls."""

    def build(self, workbook):
        worksheet = workbook.create_sheet("Household Setup")
        worksheet.sheet_view.showGridLines = False
        for col, width in zip(["A", "B", "C", "D", "E"], [3, 35, 24, 24, 3]):
            worksheet.column_dimensions[col].width = width
        self.title(
            worksheet,
            "HOUSEHOLD SETUP",
            "Editable assumptions control partner budgets and the shared-cost allocation.",
            end_col=4,
        )
        headers = ["Setting", "Partner 1", "Partner 2"]
        for col, header in enumerate(headers, start=2):
            cell = worksheet.cell(row=6, column=col)
            cell.value = header
            cell.font = self.styles.header_font()
            cell.fill = self.styles.header_fill()
            cell.border = self.styles.thin_border
            cell.alignment = self.styles.center_alignment()
        couple = self.config.couple
        rows = [
            ("Name", couple.partner_one.name, couple.partner_two.name, "General"),
            (
                "Monthly gross income",
                couple.partner_one.monthly_gross_income,
                couple.partner_two.monthly_gross_income,
                self.config.currency_format,
            ),
            (
                "Monthly net income",
                couple.partner_one.monthly_net_income,
                couple.partner_two.monthly_net_income,
                self.config.currency_format,
            ),
            (
                "Opening personal savings",
                couple.partner_one.opening_savings_balance,
                couple.partner_two.opening_savings_balance,
                self.config.currency_format,
            ),
        ]
        for row, (label, value_one, value_two, number_format) in enumerate(rows, start=7):
            worksheet.cell(row=row, column=2).value = label
            for col, value in ((3, value_one), (4, value_two)):
                cell = worksheet.cell(row=row, column=col)
                cell.value = value
                cell.font = self.styles.input_font()
                cell.number_format = number_format
                cell.alignment = self.styles.right_alignment() if number_format != "General" else self.styles.left_alignment()
            self.apply_border_range(worksheet, row, row, 2, 4)

        worksheet["B13"] = "Shared-cost allocation"
        worksheet["B13"].font = Font(size=12, bold=True, color=self.styles.palette.header_dark)
        allocation_rows = [
            (14, "Method", couple.contribution_method.value),
            (15, "Partner 1 custom share", couple.partner_one.custom_contribution_share or 0.5),
            (16, "Partner 2 custom share", couple.partner_two.custom_contribution_share or 0.5),
            (17, "Partner 1 active share", self._active_share_formula(partner_one=True)),
            (18, "Partner 2 active share", "=1-C17"),
            (19, "Allocation check", "=C17+C18"),
            (20, "Opening joint savings", couple.opening_joint_savings_balance),
            (21, "Emergency-fund months", couple.emergency_fund_months),
        ]
        for row, label, value in allocation_rows:
            worksheet[f"B{row}"] = label
            worksheet[f"C{row}"] = value
            worksheet[f"C{row}"].font = self.styles.input_font() if row in (14, 15, 16, 20, 21) else self.styles.calc_font()
            worksheet[f"C{row}"].alignment = self.styles.right_alignment()
            worksheet[f"C{row}"].number_format = self.config.percent_format if row in (15, 16, 17, 18, 19) else (
                self.config.currency_format if row == 20 else "General"
            )
            self.apply_border_range(worksheet, row, row, 2, 3)
        validation = DataValidation(
            type="list",
            formula1='"Equal,Income proportional,Custom"',
            allow_blank=False,
        )
        worksheet.add_data_validation(validation)
        validation.add(worksheet["C14"])
        share_validation = DataValidation(
            type="decimal",
            operator="between",
            formula1="0",
            formula2="1",
            allow_blank=False,
        )
        worksheet.add_data_validation(share_validation)
        share_validation.add("C15:C16")
        red_fill = PatternFill(
            start_color=self.styles.palette.red_fill,
            end_color=self.styles.palette.red_fill,
            fill_type="solid",
        )
        worksheet.conditional_formatting.add(
            "C19",
            CellIsRule(operator="notEqual", formula=["1"], fill=red_fill),
        )
        worksheet["B24"] = "Contribution shares must add up to 100%."
        worksheet["B24"].font = Font(size=10, italic=True, color="666666")
        worksheet.freeze_panes = "C7"
        return worksheet

    @staticmethod
    def _active_share_formula(partner_one: bool) -> str:
        custom_cell = "C15" if partner_one else "C16"
        income_cell = "C9" if partner_one else "D9"
        return (
            f'=IF(C14="Equal",50%,IF(C14="Income proportional",'
            f"IF(C9+D9=0,50%,{income_cell}/(C9+D9)),{custom_cell}))"
        )


class PartnerBudgetSheetBuilder(CoupleBaseSheetBuilder):
    """Build one partner's personal income, expenses and savings budget."""

    def __init__(self, config, styles, rows, partner: PersonConfig, partner_index: int) -> None:
        super().__init__(config, styles, rows)
        self.partner = partner
        self.partner_index = partner_index
        self.sheet_name = f"Partner {partner_index} Budget"
        self.table_name = f"Partner{partner_index}ExpenseTable"

    def build(self, workbook):
        worksheet = workbook.create_sheet(self.sheet_name)
        worksheet.sheet_view.showGridLines = False
        self.configure_monthly_columns(worksheet)
        worksheet.column_dimensions["C"].width = 18
        worksheet.column_dimensions["D"].width = 24
        for col in (
            self.year_budget_col() + 2,
            self.year_actual_col() + 2,
            self.year_variance_col() + 2,
        ):
            worksheet.column_dimensions[get_column_letter(col)].width = 14
        self.title(
            worksheet,
            f"{self.partner.name.upper()} PERSONAL BUDGET",
            "Personal income, expenses, savings and transfers to the shared household.",
        )
        setup_name_cell = "C7" if self.partner_index == 1 else "D7"
        worksheet["B2"] = f"=UPPER('Household Setup'!{setup_name_cell})&\" PERSONAL BUDGET\""
        self.create_month_header_grid(worksheet, title_row=5, subtitle_row=6)
        self._create_summary(worksheet)
        table_end = self._create_expense_table(worksheet, start_row=30)
        self._format_detail_variances(worksheet, 31, table_end)
        worksheet.freeze_panes = "C7"
        return worksheet

    def _create_summary(self, worksheet) -> None:
        rows = PARTNER_ROWS
        self.create_section_header(worksheet, 7, "PERSONAL INCOME")
        self._create_seeded_income_row(worksheet, rows.net_income, "Net income", setup_row=9)
        self.create_input_row(worksheet, rows.other_income, "Other income", indent=True)
        self.create_input_row(worksheet, rows.bonus, "Bonus", indent=True)

        self.create_section_header(worksheet, 12, "PERSONAL EXPENSES")
        self._create_table_rollup(worksheet, rows.fixed_expenses, "Fixed personal expenses", "Fixed")
        self._create_table_rollup(worksheet, rows.variable_expenses, "Variable personal expenses", "Variable")
        self.link_row(
            worksheet,
            rows.total_expenses,
            "Total personal expenses",
            lambda i, _m: f"={self.month_budget_letter(i)}{rows.fixed_expenses}+{self.month_budget_letter(i)}{rows.variable_expenses}",
            lambda i, _m: f"={self.month_actual_letter(i)}{rows.fixed_expenses}+{self.month_actual_letter(i)}{rows.variable_expenses}",
        )

        self.create_section_header(worksheet, 17, "SAVINGS & HOUSEHOLD TRANSFERS")
        self.create_input_row(worksheet, rows.personal_savings, "Personal savings")
        share_cell = "C17" if self.partner_index == 1 else "C18"
        self.link_row(
            worksheet,
            rows.household_contribution,
            "Contribution to shared household",
            lambda i, _m: f"='Shared Household'!{self.month_budget_letter(i)}{SHARED_ROWS.required_funding}*'Household Setup'!${share_cell}",
            lambda _i, _m: "",
        )
        for index in range(len(self.config.months)):
            cell = worksheet.cell(
                row=rows.household_contribution,
                column=self.month_actual_col(index),
            )
            cell.value = None
            cell.font = self.styles.input_font()
        self._create_opening_savings_rollforward(worksheet)

        self.create_section_header(worksheet, 22, "PERSONAL SUMMARY")
        self.link_row(
            worksheet,
            rows.total_income,
            "Total personal income",
            lambda i, _m: "+".join(
                [
                    f"={self.month_budget_letter(i)}{rows.net_income}",
                    f"{self.month_budget_letter(i)}{rows.other_income}",
                    f"{self.month_budget_letter(i)}{rows.bonus}",
                ]
            ),
            lambda i, _m: "+".join(
                [
                    f"={self.month_actual_letter(i)}{rows.net_income}",
                    f"{self.month_actual_letter(i)}{rows.other_income}",
                    f"{self.month_actual_letter(i)}{rows.bonus}",
                ]
            ),
        )
        self.link_row(
            worksheet,
            rows.disposable_balance,
            "Disposable balance",
            lambda i, _m: (
                f"={self.month_budget_letter(i)}{rows.total_income}-{self.month_budget_letter(i)}{rows.total_expenses}"
                f"-{self.month_budget_letter(i)}{rows.personal_savings}-{self.month_budget_letter(i)}{rows.household_contribution}"
            ),
            lambda i, _m: (
                f"={self.month_actual_letter(i)}{rows.total_income}-{self.month_actual_letter(i)}{rows.total_expenses}"
                f"-{self.month_actual_letter(i)}{rows.personal_savings}-{self.month_actual_letter(i)}{rows.household_contribution}"
            ),
        )
        self.link_row(
            worksheet,
            rows.savings_rate,
            "Personal savings rate",
            lambda i, _m: f"=IF({self.month_budget_letter(i)}{rows.total_income}=0,0,{self.month_budget_letter(i)}{rows.personal_savings}/{self.month_budget_letter(i)}{rows.total_income})",
            lambda i, _m: f"=IF({self.month_actual_letter(i)}{rows.total_income}=0,0,{self.month_actual_letter(i)}{rows.personal_savings}/{self.month_actual_letter(i)}{rows.total_income})",
            percent=True,
        )
        self.link_row(
            worksheet,
            rows.closing_savings,
            "Closing personal savings",
            lambda i, _m: f"={self.month_budget_letter(i)}{rows.opening_savings}+{self.month_budget_letter(i)}{rows.personal_savings}",
            lambda i, _m: f"={self.month_actual_letter(i)}{rows.opening_savings}+{self.month_actual_letter(i)}{rows.personal_savings}",
            total_last_value=True,
        )
        self.add_input_variance_rules(worksheet, rows.fixed_expenses, rows.total_expenses)

    def _create_seeded_income_row(self, worksheet, row: int, label: str, setup_row: int) -> None:
        self._set_row_label(worksheet, row, label)
        setup_col = "C" if self.partner_index == 1 else "D"
        for index in range(len(self.config.months)):
            budget = worksheet.cell(row=row, column=self.month_budget_col(index))
            actual = worksheet.cell(row=row, column=self.month_actual_col(index))
            variance = worksheet.cell(row=row, column=self.month_variance_col(index))
            budget.value = f"='Household Setup'!${setup_col}${setup_row}"
            actual.value = None
            variance.value = f"={self.month_actual_letter(index)}{row}-{self.month_budget_letter(index)}{row}"
            for cell in (budget, actual, variance):
                cell.border = self.styles.thin_border
                cell.number_format = self.config.currency_format
                cell.alignment = self.styles.right_alignment()
            budget.font = self.styles.calc_font()
            actual.font = self.styles.input_font()
            variance.font = self.styles.calc_font()
        self._set_year_total_formulas(worksheet, row, use_percent=False)

    def _create_table_rollup(self, worksheet, row: int, label: str, expense_type: str) -> None:
        self.link_row(
            worksheet,
            row,
            label,
            lambda _i, month: f'=SUMIFS({self.table_name}[{month} Budget],{self.table_name}[Type],"{expense_type}")',
            lambda _i, month: f'=SUMIFS({self.table_name}[{month} Actual],{self.table_name}[Type],"{expense_type}")',
        )

    def _create_opening_savings_rollforward(self, worksheet) -> None:
        row = PARTNER_ROWS.opening_savings
        self._set_row_label(worksheet, row, "Opening personal savings")
        setup_col = "C" if self.partner_index == 1 else "D"
        for index in range(len(self.config.months)):
            if index == 0:
                budget_formula = f"='Household Setup'!${setup_col}$10"
                actual_formula = f"='Household Setup'!${setup_col}$10"
            else:
                budget_formula = f"={self.month_budget_letter(index - 1)}{PARTNER_ROWS.closing_savings}"
                actual_formula = f"={self.month_actual_letter(index - 1)}{PARTNER_ROWS.closing_savings}"
            for col, formula in (
                (self.month_budget_col(index), budget_formula),
                (self.month_actual_col(index), actual_formula),
            ):
                cell = worksheet.cell(row=row, column=col)
                cell.value = formula
                cell.font = self.styles.calc_font()
                cell.border = self.styles.thin_border
                cell.number_format = self.config.currency_format
            variance = worksheet.cell(row=row, column=self.month_variance_col(index))
            variance.value = f"={self.month_actual_letter(index)}{row}-{self.month_budget_letter(index)}{row}"
            variance.font = self.styles.calc_font()
            variance.border = self.styles.thin_border
            variance.number_format = self.config.currency_format
        self.link_row(
            worksheet,
            row,
            "Opening personal savings",
            lambda i, _m: worksheet.cell(row=row, column=self.month_budget_col(i)).value,
            lambda i, _m: worksheet.cell(row=row, column=self.month_actual_col(i)).value,
            total_last_value=True,
        )

    def _create_expense_table(self, worksheet, start_row: int) -> int:
        self.create_section_header(worksheet, start_row - 2, "PERSONAL EXPENSE DETAIL")
        headers = ["Type", "Category", "Item", *self._detail_headers(), "Year Budget", "Year Actual", "Year Variance"]
        for col, header in enumerate(headers, start=2):
            cell = worksheet.cell(row=start_row, column=col)
            cell.value = header
            cell.font = self.styles.header_font()
            cell.fill = self.styles.header_fill()
            cell.border = self.styles.thin_border
            cell.alignment = self.styles.center_alignment()
        seeds = [
            ("Fixed", "Connectivity", "Data / mobile"),
            ("Fixed", "Transport", "Transport"),
            ("Fixed", "Subscriptions", "Subscriptions"),
            ("Fixed", "Debt", "Debt repayment"),
            ("Fixed", "Family", "Family support"),
            ("Variable", "Personal care", "Personal care"),
            ("Variable", "Dining", "Eating out"),
            ("Variable", "Clothing", "Clothing"),
            ("Variable", "Entertainment", "Entertainment"),
            ("Variable", "Other", "Other personal spending"),
        ]
        row = start_row + 1
        for expense_type, category, item in seeds:
            self._write_expense_row(worksheet, row, expense_type, category, item)
            row += 1
        for _ in range(5):
            self._write_expense_row(worksheet, row, "Variable", "Other", "")
            row += 1
        end_col = self.year_variance_col() + 2
        table = Table(displayName=self.table_name, ref=f"B{start_row}:{get_column_letter(end_col)}{row - 1}")
        table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
        worksheet.add_table(table)
        return row - 1

    def _detail_headers(self):
        for month in self.config.months:
            yield f"{month} Budget"
            yield f"{month} Actual"
            yield f"{month} Variance"

    def _write_expense_row(self, worksheet, row: int, expense_type: str, category: str, item: str) -> None:
        for col, value in ((2, expense_type), (3, category), (4, item)):
            worksheet.cell(row=row, column=col).value = value
            worksheet.cell(row=row, column=col).border = self.styles.thin_border
        for index in range(len(self.config.months)):
            budget_col = self.month_budget_col(index) + 2
            actual_col = self.month_actual_col(index) + 2
            variance_col = self.month_variance_col(index) + 2
            worksheet.cell(row=row, column=variance_col).value = (
                f"={get_column_letter(actual_col)}{row}-{get_column_letter(budget_col)}{row}"
            )
            for col in (budget_col, actual_col, variance_col):
                cell = worksheet.cell(row=row, column=col)
                cell.border = self.styles.thin_border
                cell.number_format = self.config.currency_format
                cell.font = self.styles.input_font() if col != variance_col else self.styles.calc_font()
        budget_cells = [f"{get_column_letter(self.month_budget_col(i) + 2)}{row}" for i in range(len(self.config.months))]
        actual_cells = [f"{get_column_letter(self.month_actual_col(i) + 2)}{row}" for i in range(len(self.config.months))]
        year_budget_col = self.year_budget_col() + 2
        year_actual_col = self.year_actual_col() + 2
        year_variance_col = self.year_variance_col() + 2
        worksheet.cell(row=row, column=year_budget_col).value = f"=SUM({','.join(budget_cells)})"
        worksheet.cell(row=row, column=year_actual_col).value = f"=SUM({','.join(actual_cells)})"
        worksheet.cell(row=row, column=year_variance_col).value = (
            f"={get_column_letter(year_actual_col)}{row}-{get_column_letter(year_budget_col)}{row}"
        )
        for col in (year_budget_col, year_actual_col, year_variance_col):
            cell = worksheet.cell(row=row, column=col)
            cell.border = self.styles.thin_border
            cell.number_format = self.config.currency_format
            cell.font = self.styles.calc_font()

    def _format_detail_variances(self, worksheet, start_row: int, end_row: int) -> None:
        columns = [self.month_variance_col(i) + 2 for i in range(len(self.config.months))]
        columns.append(self.year_variance_col() + 2)
        self.add_variance_conditional_formatting(worksheet, start_row, end_row, columns=columns)


class SharedHouseholdSheetBuilder(PartnerBudgetSheetBuilder):
    """Build shared household expenses, contributions and joint savings."""

    def __init__(self, config, styles, rows) -> None:
        super().__init__(config, styles, rows, config.couple.partner_one, 1)
        self.sheet_name = "Shared Household"
        self.table_name = "SharedExpenseTable"

    def build(self, workbook):
        worksheet = workbook.create_sheet(self.sheet_name)
        worksheet.sheet_view.showGridLines = False
        self.configure_monthly_columns(worksheet)
        self.title(
            worksheet,
            "SHARED HOUSEHOLD BUDGET",
            "Shared expenses, joint savings and contributions received from both partners.",
        )
        self.create_month_header_grid(worksheet, title_row=5, subtitle_row=6)
        self._create_shared_summary(worksheet)
        table_end = self._create_shared_expense_table(worksheet, start_row=28)
        self._format_detail_variances(worksheet, 29, table_end)
        worksheet.freeze_panes = "C7"
        return worksheet

    def _create_shared_summary(self, worksheet) -> None:
        rows = SHARED_ROWS
        partner_rows = PARTNER_ROWS
        self.create_section_header(worksheet, 7, "CONTRIBUTIONS RECEIVED")
        for row, sheet_name, name in (
            (rows.partner_one_contribution, "Partner 1 Budget", self.config.couple.partner_one.name),
            (rows.partner_two_contribution, "Partner 2 Budget", self.config.couple.partner_two.name),
        ):
            self.link_row(
                worksheet,
                row,
                f"{name} contribution",
                lambda i, _m, source=sheet_name: f"='{source}'!{self.month_budget_letter(i)}{partner_rows.household_contribution}",
                lambda i, _m, source=sheet_name: f"='{source}'!{self.month_actual_letter(i)}{partner_rows.household_contribution}",
            )
            setup_name_cell = "C7" if sheet_name == "Partner 1 Budget" else "D7"
            worksheet.cell(row=row, column=2).value = (
                f"='Household Setup'!{setup_name_cell}&\" contribution\""
            )
        self.link_row(
            worksheet,
            rows.total_contributions,
            "Total contributions",
            lambda i, _m: f"={self.month_budget_letter(i)}{rows.partner_one_contribution}+{self.month_budget_letter(i)}{rows.partner_two_contribution}",
            lambda i, _m: f"={self.month_actual_letter(i)}{rows.partner_one_contribution}+{self.month_actual_letter(i)}{rows.partner_two_contribution}",
        )
        self.create_section_header(worksheet, 12, "SHARED EXPENSES")
        self._create_table_rollup(worksheet, rows.fixed_expenses, "Fixed shared expenses", "Fixed")
        self._create_table_rollup(worksheet, rows.variable_expenses, "Variable shared expenses", "Variable")
        self.link_row(
            worksheet,
            rows.total_expenses,
            "Total shared expenses",
            lambda i, _m: f"={self.month_budget_letter(i)}{rows.fixed_expenses}+{self.month_budget_letter(i)}{rows.variable_expenses}",
            lambda i, _m: f"={self.month_actual_letter(i)}{rows.fixed_expenses}+{self.month_actual_letter(i)}{rows.variable_expenses}",
        )
        self.create_section_header(worksheet, 17, "JOINT SAVINGS")
        self.create_input_row(worksheet, rows.joint_savings, "Transfer to joint savings")
        self._create_joint_savings_opening(worksheet)
        self.create_section_header(worksheet, 21, "SHARED SUMMARY")
        self.link_row(
            worksheet,
            rows.required_funding,
            "Required shared funding",
            lambda i, _m: f"={self.month_budget_letter(i)}{rows.total_expenses}+{self.month_budget_letter(i)}{rows.joint_savings}",
            lambda i, _m: f"={self.month_actual_letter(i)}{rows.total_expenses}+{self.month_actual_letter(i)}{rows.joint_savings}",
        )
        self.link_row(
            worksheet,
            rows.shared_balance,
            "Shared account surplus / (deficit)",
            lambda i, _m: f"={self.month_budget_letter(i)}{rows.total_contributions}-{self.month_budget_letter(i)}{rows.required_funding}",
            lambda i, _m: f"={self.month_actual_letter(i)}{rows.total_contributions}-{self.month_actual_letter(i)}{rows.required_funding}",
        )
        self.link_row(
            worksheet,
            rows.closing_joint_savings,
            "Closing joint savings",
            lambda i, _m: f"={self.month_budget_letter(i)}{rows.opening_joint_savings}+{self.month_budget_letter(i)}{rows.joint_savings}",
            lambda i, _m: f"={self.month_actual_letter(i)}{rows.opening_joint_savings}+{self.month_actual_letter(i)}{rows.joint_savings}",
            total_last_value=True,
        )
        self.add_input_variance_rules(worksheet, rows.fixed_expenses, rows.total_expenses)

    def _create_joint_savings_opening(self, worksheet) -> None:
        row = SHARED_ROWS.opening_joint_savings
        self._set_row_label(worksheet, row, "Opening joint savings")
        for index in range(len(self.config.months)):
            if index == 0:
                budget_formula = actual_formula = "='Household Setup'!$C$20"
            else:
                budget_formula = f"={self.month_budget_letter(index - 1)}{SHARED_ROWS.closing_joint_savings}"
                actual_formula = f"={self.month_actual_letter(index - 1)}{SHARED_ROWS.closing_joint_savings}"
            for col, formula in (
                (self.month_budget_col(index), budget_formula),
                (self.month_actual_col(index), actual_formula),
            ):
                cell = worksheet.cell(row=row, column=col)
                cell.value = formula
                cell.font = self.styles.calc_font()
                cell.border = self.styles.thin_border
                cell.number_format = self.config.currency_format
            variance = worksheet.cell(row=row, column=self.month_variance_col(index))
            variance.value = f"={self.month_actual_letter(index)}{row}-{self.month_budget_letter(index)}{row}"
            variance.font = self.styles.calc_font()
            variance.border = self.styles.thin_border
            variance.number_format = self.config.currency_format
        self.link_row(
            worksheet,
            row,
            "Opening joint savings",
            lambda i, _m: worksheet.cell(row=row, column=self.month_budget_col(i)).value,
            lambda i, _m: worksheet.cell(row=row, column=self.month_actual_col(i)).value,
            total_last_value=True,
        )

    def _create_shared_expense_table(self, worksheet, start_row: int) -> int:
        self.create_section_header(worksheet, start_row - 2, "SHARED EXPENSE DETAIL")
        headers = ["Type", "Category", "Item", *self._detail_headers(), "Year Budget", "Year Actual", "Year Variance"]
        for col, header in enumerate(headers, start=2):
            cell = worksheet.cell(row=start_row, column=col)
            cell.value = header
            cell.font = self.styles.header_font()
            cell.fill = self.styles.header_fill()
            cell.border = self.styles.thin_border
            cell.alignment = self.styles.center_alignment()
        seeds = [
            ("Fixed", "Housing", "Rent / bond", self.config.default_rent_amount),
            ("Fixed", "Utilities", "Electricity", 0),
            ("Fixed", "Utilities", "Water", 0),
            ("Fixed", "Connectivity", "Internet", 0),
            ("Fixed", "Subscriptions", "Shared subscriptions", 0),
            ("Variable", "Groceries", "Groceries", 0),
            ("Variable", "Household", "Household supplies", 0),
            ("Variable", "Transport", "Shared transport", 0),
            ("Variable", "Entertainment", "Dates / entertainment", 0),
            ("Variable", "Other", "Other shared expenses", 0),
        ]
        row = start_row + 1
        for expense_type, category, item, monthly_budget in seeds:
            self._write_expense_row(worksheet, row, expense_type, category, item)
            if monthly_budget:
                for index in range(len(self.config.months)):
                    worksheet.cell(row=row, column=self.month_budget_col(index) + 2).value = monthly_budget
            row += 1
        for _ in range(5):
            self._write_expense_row(worksheet, row, "Variable", "Other", "")
            row += 1
        end_col = self.year_variance_col() + 2
        table = Table(displayName=self.table_name, ref=f"B{start_row}:{get_column_letter(end_col)}{row - 1}")
        table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
        worksheet.add_table(table)
        return row - 1


class HouseholdDashboardSheetBuilder(CoupleBaseSheetBuilder):
    """Build a consolidated dashboard without double-counting internal transfers."""

    metrics = (
        ("Combined external income", PARTNER_ROWS.total_income, PARTNER_ROWS.total_income, None),
        ("Personal expenses", PARTNER_ROWS.total_expenses, PARTNER_ROWS.total_expenses, None),
        ("Shared expenses", None, None, SHARED_ROWS.total_expenses),
        ("Total external expenses", "expenses", "expenses", SHARED_ROWS.total_expenses),
        ("Personal savings", PARTNER_ROWS.personal_savings, PARTNER_ROWS.personal_savings, None),
        ("Joint savings", None, None, SHARED_ROWS.joint_savings),
        ("Total savings", "savings", "savings", SHARED_ROWS.joint_savings),
        ("Household surplus / (deficit)", "surplus", "surplus", None),
        ("Household savings rate", "rate", "rate", None),
    )

    def build(self, workbook):
        worksheet = workbook.create_sheet("Household Dashboard", 1)
        worksheet.sheet_view.showGridLines = False
        self.configure_monthly_columns(worksheet, label_width=32)
        self.title(
            worksheet,
            "HOUSEHOLD DASHBOARD",
            "Combined results exclude partner contributions because they are internal transfers.",
        )
        self.create_month_header_grid(worksheet, title_row=5, subtitle_row=6)
        self.create_section_header(worksheet, 7, "CONSOLIDATED HOUSEHOLD RESULTS")
        start_row = 8
        for offset, (label, p1_key, p2_key, shared_row) in enumerate(self.metrics):
            row = start_row + offset
            self.link_row(
                worksheet,
                row,
                label,
                lambda i, _m, key=p1_key, srow=shared_row: self._metric_formula(i, key, srow, actual=False),
                lambda i, _m, key=p2_key, srow=shared_row: self._metric_formula(i, key, srow, actual=True),
                percent=label == "Household savings rate",
            )
        savings_rate_row = start_row + len(self.metrics) - 1
        worksheet.cell(row=savings_rate_row, column=self.year_budget_col()).value = (
            f"=IF({self.year_budget_letter()}{start_row}=0,0,"
            f"{self.year_budget_letter()}{start_row + 6}/{self.year_budget_letter()}{start_row})"
        )
        worksheet.cell(row=savings_rate_row, column=self.year_actual_col()).value = (
            f"=IF({self.year_actual_letter()}{start_row}=0,0,"
            f"{self.year_actual_letter()}{start_row + 6}/{self.year_actual_letter()}{start_row})"
        )
        worksheet.cell(row=savings_rate_row, column=self.year_variance_col()).value = (
            f"={self.year_actual_letter()}{savings_rate_row}-"
            f"{self.year_budget_letter()}{savings_rate_row}"
        )
        self._create_contribution_review(worksheet, start_row=20)
        self._add_charts(worksheet, start_row)
        worksheet.freeze_panes = "C7"
        return worksheet

    def _metric_formula(self, index: int, key, shared_row: int | None, *, actual: bool) -> str:
        col = self.month_actual_letter(index) if actual else self.month_budget_letter(index)
        p1 = "'Partner 1 Budget'"
        p2 = "'Partner 2 Budget'"
        shared = "'Shared Household'"
        if isinstance(key, int):
            return f"={p1}!{col}{key}+{p2}!{col}{key}"
        if key is None and shared_row is not None:
            return f"={shared}!{col}{shared_row}"
        if key == "expenses":
            return (
                f"={p1}!{col}{PARTNER_ROWS.total_expenses}+{p2}!{col}{PARTNER_ROWS.total_expenses}"
                f"+{shared}!{col}{shared_row}"
            )
        if key == "savings":
            return (
                f"={p1}!{col}{PARTNER_ROWS.personal_savings}+{p2}!{col}{PARTNER_ROWS.personal_savings}"
                f"+{shared}!{col}{shared_row}"
            )
        income_row = 8
        expense_row = 11
        savings_row = 14
        if key == "surplus":
            return f"={col}{income_row}-{col}{expense_row}-{col}{savings_row}"
        if key == "rate":
            return f"=IF({col}{income_row}=0,0,{col}{savings_row}/{col}{income_row})"
        raise ValueError(f"Unsupported household metric key: {key}")

    def _create_contribution_review(self, worksheet, start_row: int) -> None:
        worksheet[f"B{start_row}"] = "SHARED CONTRIBUTION REVIEW"
        worksheet[f"B{start_row}"].font = Font(size=14, bold=True, color=self.styles.palette.header_dark)
        headers = ["Partner", "Active share", "Expected annual", "Actual annual", "Difference"]
        for col, header in enumerate(headers, start=2):
            cell = worksheet.cell(row=start_row + 2, column=col)
            cell.value = header
            cell.font = self.styles.header_font()
            cell.fill = self.styles.header_fill()
            cell.border = self.styles.thin_border
        for offset, (name, setup_share, source_sheet) in enumerate(
            (
                (self.config.couple.partner_one.name, "C17", "Partner 1 Budget"),
                (self.config.couple.partner_two.name, "C18", "Partner 2 Budget"),
            ),
            start=3,
        ):
            row = start_row + offset
            setup_name_cell = "C7" if source_sheet == "Partner 1 Budget" else "D7"
            worksheet[f"B{row}"] = f"='Household Setup'!{setup_name_cell}"
            worksheet[f"C{row}"] = f"='Household Setup'!{setup_share}"
            worksheet[f"D{row}"] = f"='{source_sheet}'!{self.year_budget_letter()}{PARTNER_ROWS.household_contribution}"
            worksheet[f"E{row}"] = f"='{source_sheet}'!{self.year_actual_letter()}{PARTNER_ROWS.household_contribution}"
            worksheet[f"F{row}"] = f"=E{row}-D{row}"
            for col in range(2, 7):
                worksheet.cell(row=row, column=col).border = self.styles.thin_border
            worksheet[f"C{row}"].number_format = self.config.percent_format
            for col in ("D", "E", "F"):
                worksheet[f"{col}{row}"].number_format = self.config.currency_format

    def _add_charts(self, worksheet, start_row: int) -> None:
        expense_chart = BarChart()
        expense_chart.title = "Annual household budget vs actual"
        expense_chart.y_axis.title = "Amount (R)"
        data = Reference(
            worksheet,
            min_col=self.year_budget_col(),
            max_col=self.year_actual_col(),
            min_row=start_row,
            max_row=start_row + 6,
        )
        categories = Reference(worksheet, min_col=2, min_row=start_row, max_row=start_row + 6)
        expense_chart.add_data(data, titles_from_data=False)
        expense_chart.series[0].title = SeriesLabel(v="Budget")
        expense_chart.series[1].title = SeriesLabel(v="Actual")
        expense_chart.set_categories(categories)
        expense_chart.height = 8
        expense_chart.width = 16
        worksheet.add_chart(expense_chart, "H20")


class HouseholdTrendSheetBuilder(CoupleBaseSheetBuilder):
    """Build a compact month-by-month household trend table and chart."""

    def build(self, workbook):
        worksheet = workbook.create_sheet("Trend Analysis")
        worksheet.sheet_view.showGridLines = False
        for col, width in zip(["A", "B", "C", "D", "E", "F"], [3, 14, 18, 18, 18, 18]):
            worksheet.column_dimensions[col].width = width
        self.title(worksheet, "HOUSEHOLD TREND ANALYSIS", "Monthly actual household results.", end_col=6)
        headers = ["Month", "Income", "Expenses", "Savings", "Surplus / (Deficit)"]
        for col, header in enumerate(headers, start=2):
            cell = worksheet.cell(row=6, column=col)
            cell.value = header
            cell.font = self.styles.header_font()
            cell.fill = self.styles.header_fill()
            cell.border = self.styles.thin_border
        dashboard_rows = (8, 11, 14, 15)
        for index, month in enumerate(self.config.months):
            row = 7 + index
            worksheet[f"B{row}"] = month
            for col, source_row in zip(("C", "D", "E", "F"), dashboard_rows):
                worksheet[f"{col}{row}"] = f"='Household Dashboard'!{self.month_actual_letter(index)}{source_row}"
                worksheet[f"{col}{row}"].number_format = self.config.currency_format
                worksheet[f"{col}{row}"].font = self.styles.calc_font()
            self.apply_border_range(worksheet, row, row, 2, 6)
        chart = LineChart()
        chart.title = "Monthly household cash flow"
        chart.y_axis.title = "Amount (R)"
        chart.x_axis.title = "Month"
        final_row = 6 + len(self.config.months)
        chart.add_data(
            Reference(worksheet, min_col=3, max_col=6, min_row=6, max_row=final_row),
            titles_from_data=True,
        )
        for series, title in zip(
            chart.series,
            ("Income", "Expenses", "Savings", "Surplus / (Deficit)"),
        ):
            series.title = SeriesLabel(v=title)
        chart.set_categories(Reference(worksheet, min_col=2, min_row=7, max_row=final_row))
        chart.height = 9
        chart.width = 18
        worksheet.add_chart(chart, "H5")
        return worksheet


class HouseholdSavingsSheetBuilder(CoupleBaseSheetBuilder):
    """Build combined personal/joint savings and emergency-fund progress."""

    def build(self, workbook):
        worksheet = workbook.create_sheet("Savings Goals")
        worksheet.sheet_view.showGridLines = False
        for col, width in zip(["A", "B", "C", "D", "E"], [3, 34, 18, 18, 18]):
            worksheet.column_dimensions[col].width = width
        self.title(
            worksheet,
            "HOUSEHOLD SAVINGS GOALS",
            "Combined personal and joint savings compared with the emergency-fund target.",
            end_col=5,
        )
        for col, header in enumerate(["Metric", "Budget", "Actual", "Variance"], start=2):
            cell = worksheet.cell(row=6, column=col)
            cell.value = header
            cell.font = self.styles.header_font()
            cell.fill = self.styles.header_fill()
            cell.border = self.styles.thin_border
        metrics = [
            (
                "Average monthly external expenses",
                f"='Household Dashboard'!{self.year_budget_letter()}11/{len(self.config.months)}",
                f"='Household Dashboard'!{self.year_actual_letter()}11/{len(self.config.months)}",
                self.config.currency_format,
            ),
            (
                "Emergency-fund target",
                "=C7*'Household Setup'!C21",
                "=D7*'Household Setup'!C21",
                self.config.currency_format,
            ),
            (
                "Combined closing savings",
                f"='Partner 1 Budget'!{self.year_budget_letter()}{PARTNER_ROWS.closing_savings}+'Partner 2 Budget'!{self.year_budget_letter()}{PARTNER_ROWS.closing_savings}+'Shared Household'!{self.year_budget_letter()}{SHARED_ROWS.closing_joint_savings}",
                f"='Partner 1 Budget'!{self.year_actual_letter()}{PARTNER_ROWS.closing_savings}+'Partner 2 Budget'!{self.year_actual_letter()}{PARTNER_ROWS.closing_savings}+'Shared Household'!{self.year_actual_letter()}{SHARED_ROWS.closing_joint_savings}",
                self.config.currency_format,
            ),
            ("Emergency-fund progress", "=IF(C8=0,0,C9/C8)", "=IF(D8=0,0,D9/D8)", self.config.percent_format),
            ("Amount still needed", "=MAX(0,C8-C9)", "=MAX(0,D8-D9)", self.config.currency_format),
        ]
        for row, (label, budget, actual, number_format) in enumerate(metrics, start=7):
            worksheet[f"B{row}"] = label
            worksheet[f"C{row}"] = budget
            worksheet[f"D{row}"] = actual
            worksheet[f"E{row}"] = f"=D{row}-C{row}"
            for col in range(2, 6):
                cell = worksheet.cell(row=row, column=col)
                cell.border = self.styles.thin_border
                if col > 2:
                    cell.number_format = number_format
                    cell.font = self.styles.calc_font()
        return worksheet


class HouseholdBonusTrackerSheetBuilder(CoupleBaseSheetBuilder):
    """Build combined and partner-level monthly bonus tracking."""

    def build(self, workbook):
        worksheet = workbook.create_sheet("Bonus Tracker")
        worksheet.sheet_view.showGridLines = False
        for col, width in zip(["A", "B", "C", "D", "E"], [3, 14, 18, 18, 18]):
            worksheet.column_dimensions[col].width = width
        self.title(worksheet, "HOUSEHOLD BONUS TRACKER", "Monthly bonuses by partner and combined.", end_col=5)
        headers = ["Month", self.config.couple.partner_one.name, self.config.couple.partner_two.name, "Combined"]
        for col, header in enumerate(headers, start=2):
            cell = worksheet.cell(row=6, column=col)
            cell.value = header
            cell.font = self.styles.header_font()
            cell.fill = self.styles.header_fill()
            cell.border = self.styles.thin_border
        worksheet["C6"] = "='Household Setup'!C7"
        worksheet["D6"] = "='Household Setup'!D7"
        for index, month in enumerate(self.config.months):
            row = 7 + index
            actual_col = self.month_actual_letter(index)
            worksheet[f"B{row}"] = month
            worksheet[f"C{row}"] = f"='Partner 1 Budget'!{actual_col}{PARTNER_ROWS.bonus}"
            worksheet[f"D{row}"] = f"='Partner 2 Budget'!{actual_col}{PARTNER_ROWS.bonus}"
            worksheet[f"E{row}"] = f"=C{row}+D{row}"
            for col in range(2, 6):
                cell = worksheet.cell(row=row, column=col)
                cell.border = self.styles.thin_border
                if col > 2:
                    cell.number_format = self.config.currency_format
                    cell.font = self.styles.calc_font()
        return worksheet


class HouseholdProjectionSheetBuilder(CoupleBaseSheetBuilder):
    """Build a five-year projection from total household savings."""

    def build(self, workbook):
        worksheet = workbook.create_sheet("5-Year Projection")
        worksheet.sheet_view.showGridLines = False
        for col, width in zip(["A", "B", "C", "D", "E", "F"], [3, 32, 20, 20, 20, 20]):
            worksheet.column_dimensions[col].width = width
        self.title(
            worksheet,
            "5-YEAR HOUSEHOLD SAVINGS PROJECTION",
            "Projection uses combined personal and joint savings plus the configured growth rate.",
            end_col=6,
        )
        worksheet["B6"] = "Annual growth rate"
        worksheet["C6"] = self.config.annual_growth_rate
        worksheet["C6"].font = self.styles.input_font()
        worksheet["C6"].number_format = self.config.percent_format
        worksheet["B7"] = "Annual budget contribution"
        worksheet["C7"] = f"='Household Dashboard'!{self.year_budget_letter()}14"
        worksheet["B8"] = "Annual actual contribution"
        worksheet["C8"] = f"='Household Dashboard'!{self.year_actual_letter()}14"
        worksheet["B9"] = "Starting budget balance"
        worksheet["C9"] = "='Savings Goals'!C9"
        worksheet["B10"] = "Starting actual balance"
        worksheet["C10"] = "='Savings Goals'!D9"
        for row in range(6, 11):
            self.apply_border_range(worksheet, row, row, 2, 3)
            if row > 6:
                worksheet[f"C{row}"].font = self.styles.calc_font()
                worksheet[f"C{row}"].number_format = self.config.currency_format
        headers = ["Year", "Budget start", "Actual start", "Budget end", "Actual end"]
        for col, header in enumerate(headers, start=2):
            cell = worksheet.cell(row=13, column=col)
            cell.value = header
            cell.font = self.styles.header_font()
            cell.fill = self.styles.header_fill()
            cell.border = self.styles.thin_border
        for year in range(1, 6):
            row = 13 + year
            worksheet[f"B{row}"] = f"Year {year}"
            worksheet[f"C{row}"] = "=$C$9" if year == 1 else f"=E{row - 1}"
            worksheet[f"D{row}"] = "=$C$10" if year == 1 else f"=F{row - 1}"
            worksheet[f"E{row}"] = f"=C{row}+$C$7+(C{row}+$C$7/2)*$C$6"
            worksheet[f"F{row}"] = f"=D{row}+$C$8+(D{row}+$C$8/2)*$C$6"
            for col in range(2, 7):
                cell = worksheet.cell(row=row, column=col)
                cell.border = self.styles.thin_border
                if col > 2:
                    cell.number_format = self.config.currency_format
                    cell.font = self.styles.calc_font()
        return worksheet
