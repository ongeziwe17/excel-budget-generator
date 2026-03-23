"""Cover worksheet builder."""

from openpyxl.styles import Alignment, Font, PatternFill

from .base import BaseSheetBuilder


class CoverSheetBuilder(BaseSheetBuilder):
    """Build the workbook cover sheet."""

    def build(self, workbook):
        worksheet = workbook.active
        worksheet.title = "Cover"
        worksheet.sheet_view.showGridLines = False

        worksheet.column_dimensions["A"].width = 3
        worksheet.column_dimensions["B"].width = 35
        worksheet.column_dimensions["C"].width = 50
        worksheet.column_dimensions["D"].width = 3

        worksheet.merge_cells("B2:C2")
        worksheet["B2"] = "PERSONAL BUDGET WORKBOOK"
        worksheet["B2"].font = Font(size=24, bold=True, color=self.styles.palette.header_dark)
        worksheet["B2"].alignment = self.styles.center_alignment()
        worksheet.row_dimensions[2].height = 40

        worksheet.merge_cells("B3:C3")
        worksheet["B3"] = "South African Rand (ZAR) | Manual Entry System"
        worksheet["B3"].font = Font(size=12, italic=True, color="666666")
        worksheet["B3"].alignment = self.styles.center_alignment()

        worksheet.merge_cells("B5:C5")
        worksheet["B5"] = (
            "Track your actual income, expenses, and savings with dynamic calculations."
        )
        worksheet["B5"].font = Font(size=11, color="444444")
        worksheet["B5"].alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

        worksheet["B7"] = "KEY FEATURES"
        worksheet["B7"].font = Font(size=14, bold=True, color=self.styles.palette.header_dark)
        worksheet.row_dimensions[7].height = 25

        features = [
            ("✓ Manual Monthly Entry", "Enter your actual income and expenses each month"),
            ("✓ Automatic Calculations", "All totals, variances, and rates calculate automatically"),
            ("✓ Trend Analysis", "Month-to-month spending and savings comparisons"),
            ("✓ Emergency Fund Tracking", "Monitor your 6-month expense buffer progress"),
            ("✓ 5-Year Projection", "Long-term savings growth at 8% annual return"),
        ]

        row = 9
        for feature, description in features:
            worksheet[f"B{row}"] = feature
            worksheet[f"B{row}"].font = Font(
                size=11,
                bold=True,
                color=self.styles.palette.header_light,
            )
            worksheet[f"C{row}"] = description
            worksheet[f"C{row}"].font = Font(size=10, color="555555")
            row += 1

        worksheet[f"B{row + 1}"] = "WORKBOOK STRUCTURE"
        worksheet[f"B{row + 1}"].font = Font(
            size=14,
            bold=True,
            color=self.styles.palette.header_dark,
        )

        sheet_index = [
            ("Monthly Entry", "Input all income, expenses, savings, and bonuses"),
            ("Summary Dashboard", "Overview with expense breakdown and key metrics"),
            ("Trend Analysis", "Month-to-month changes with conditional formatting"),
            ("Bonus Tracker", "Track quarterly and annual bonus income"),
            ("Emergency Fund", "Monitor emergency fund progress vs 6-month target"),
            ("5-Year Projection", "Long-term savings projection at 8% growth"),
        ]

        row += 3
        worksheet[f"B{row}"] = "Sheet Name"
        worksheet[f"B{row}"].font = Font(bold=True, color=self.styles.palette.header_dark)
        worksheet[f"B{row}"].fill = PatternFill(
            start_color="D6E3F8",
            end_color="D6E3F8",
            fill_type="solid",
        )
        worksheet[f"C{row}"] = "Purpose"
        worksheet[f"C{row}"].font = Font(bold=True, color=self.styles.palette.header_dark)
        worksheet[f"C{row}"].fill = PatternFill(
            start_color="D6E3F8",
            end_color="D6E3F8",
            fill_type="solid",
        )
        row += 1

        for sheet_name, purpose in sheet_index:
            worksheet[f"B{row}"] = sheet_name
            worksheet[f"B{row}"].font = Font(size=10, bold=True)
            worksheet[f"C{row}"] = purpose
            worksheet[f"C{row}"].font = Font(size=10, color="555555")
            if row % 2 == 0:
                worksheet[f"B{row}"].fill = self.styles.alt_row_fill()
                worksheet[f"C{row}"].fill = self.styles.alt_row_fill()
            row += 1

        worksheet[f"B{row + 2}"] = "HOW TO USE THIS WORKBOOK"
        worksheet[f"B{row + 2}"].font = Font(
            size=14,
            bold=True,
            color=self.styles.palette.header_dark,
        )

        instructions = [
            "1. BLUE TEXT cells are for YOUR INPUT - enter actual amounts only",
            "2. BLACK TEXT cells contain FORMULAS - do not modify these",
            "3. Enter data month by month in the 'Monthly Entry' sheet",
            "4. All other sheets update automatically based on your entries",
        ]

        row += 4
        for instruction in instructions:
            worksheet[f"B{row}"] = instruction
            worksheet[f"B{row}"].font = Font(size=10, color="444444")
            row += 1

        worksheet[f"B{row + 2}"] = "COLOR LEGEND"
        worksheet[f"B{row + 2}"].font = Font(
            size=12,
            bold=True,
            color=self.styles.palette.header_dark,
        )

        row += 4
        worksheet[f"B{row}"] = "Blue Text"
        worksheet[f"B{row}"].font = Font(color=self.styles.palette.input_blue, bold=True)
        worksheet[f"C{row}"] = "Manual input cells - enter your actual amounts here"
        worksheet[f"C{row}"].font = Font(size=10)

        worksheet[f"B{row + 1}"] = "Black Text"
        worksheet[f"B{row + 1}"].font = Font(color=self.styles.palette.text_dark, bold=True)
        worksheet[f"C{row + 1}"] = "Formula cells - calculations happen automatically"
        worksheet[f"C{row + 1}"].font = Font(size=10)

        worksheet[f"B{row + 2}"] = "Green Highlight"
        worksheet[f"B{row + 2}"].fill = PatternFill(
            start_color=self.styles.palette.green_fill,
            end_color=self.styles.palette.green_fill,
            fill_type="solid",
        )
        worksheet[f"C{row + 2}"] = "Positive trend - spending decreased or savings increased"
        worksheet[f"C{row + 2}"].font = Font(size=10)

        worksheet[f"B{row + 3}"] = "Red Highlight"
        worksheet[f"B{row + 3}"].fill = PatternFill(
            start_color=self.styles.palette.red_fill,
            end_color=self.styles.palette.red_fill,
            fill_type="solid",
        )
        worksheet[f"C{row + 3}"] = "Negative trend - spending increased or savings decreased"
        worksheet[f"C{row + 3}"].font = Font(size=10)

        return worksheet
