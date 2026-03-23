"""Emergency fund worksheet builder."""

from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from .base import BaseSheetBuilder


class EmergencyFundSheetBuilder(BaseSheetBuilder):
    """Build the Emergency Fund worksheet."""

    def build(self, workbook):
        worksheet = workbook.create_sheet("Emergency Fund")
        worksheet.sheet_view.showGridLines = False

        worksheet.column_dimensions["A"].width = 3
        worksheet.column_dimensions["B"].width = 35
        worksheet.column_dimensions["C"].width = 20
        worksheet.column_dimensions["D"].width = 20
        worksheet.column_dimensions["E"].width = 20

        worksheet.merge_cells("B2:E2")
        worksheet["B2"] = "EMERGENCY FUND TRACKER"
        worksheet["B2"].font = Font(size=18, bold=True, color=self.styles.palette.header_dark)
        worksheet["B2"].alignment = self.styles.left_alignment()
        worksheet.row_dimensions[2].height = 30

        worksheet.merge_cells("B3:E3")
        worksheet["B3"] = "Monitor your progress toward a 6-month expense buffer"
        worksheet["B3"].font = Font(size=10, italic=True, color="666666")

        worksheet["B5"] = "CURRENT STATUS"
        worksheet["B5"].font = Font(size=14, bold=True, color=self.styles.palette.header_dark)
        worksheet.row_dimensions[5].height = 25

        self._create_status_metrics(worksheet)
        self._create_monthly_progress_table(worksheet)
        self._create_guidelines(worksheet)
        return worksheet

    def _create_status_metrics(self, worksheet) -> None:
        labels = [
            (7, "Average Monthly Expenses", "='Monthly Entry'!O44/12", self.styles.palette.header_light),
            (8, "6-Month Emergency Fund Target", "=C7*6", self.styles.palette.header_dark),
            (9, "Current Savings Balance", "='Monthly Entry'!O48", self.styles.palette.positive_green),
            (10, "Progress Toward Target", "=IF(C8=0,0,C9/C8)", self.styles.palette.header_light),
            (11, "Amount Still Needed", "=MAX(0,C8-C9)", self.styles.palette.negative_red),
        ]
        alt_rows = {8, 10}
        percent_rows = {10}

        for row, label, formula, color in labels:
            worksheet[f"B{row}"] = label
            worksheet[f"B{row}"].font = Font(size=11, color=self.styles.palette.text_dark)
            worksheet[f"B{row}"].border = self.styles.thin_border
            worksheet[f"B{row}"].alignment = self.styles.left_alignment()

            worksheet[f"C{row}"] = formula
            worksheet[f"C{row}"].font = Font(size=12, bold=True, color=color)
            worksheet[f"C{row}"].border = self.styles.thin_border
            worksheet[f"C{row}"].alignment = self.styles.right_alignment()
            worksheet[f"C{row}"].number_format = (
                self.config.percent_format if row in percent_rows else self.config.currency_format
            )

            if row in alt_rows:
                worksheet[f"B{row}"].fill = self.styles.alt_row_fill()
                worksheet[f"C{row}"].fill = self.styles.alt_row_fill()

        worksheet["B13"] = "EMERGENCY FUND STATUS"
        worksheet["B13"].font = Font(size=12, bold=True, color=self.styles.palette.header_dark)

        worksheet["B14"] = '=IF(C10>=1,"FULLY FUNDED",IF(C10>=0.5,"HALFWAY THERE","BUILDING..."))'
        worksheet["B14"].font = Font(size=14, bold=True)
        worksheet["B14"].alignment = self.styles.center_alignment()
        worksheet.merge_cells("B14:C14")
        worksheet.row_dimensions[14].height = 30

    def _create_monthly_progress_table(self, worksheet) -> None:
        worksheet["B17"] = "MONTHLY SAVINGS PROGRESS"
        worksheet["B17"].font = Font(size=14, bold=True, color=self.styles.palette.header_dark)

        headers = [("B19", "Month"), ("C19", "Monthly Savings"), ("D19", "Cumulative Savings"), ("E19", "% of Target")]
        for cell_ref, label in headers:
            worksheet[cell_ref] = label
            worksheet[cell_ref].font = self.styles.header_font()
            worksheet[cell_ref].fill = self.styles.header_fill()
            worksheet[cell_ref].border = self.styles.thin_border
            worksheet[cell_ref].alignment = self.styles.right_alignment() if cell_ref != "B19" else self.styles.left_alignment()
        worksheet.row_dimensions[19].height = 22

        row = 20
        for index, month in enumerate(self.config.months):
            col = get_column_letter(3 + index)
            worksheet[f"B{row}"] = month
            worksheet[f"B{row}"].font = Font(size=10, color=self.styles.palette.text_dark)
            worksheet[f"B{row}"].border = self.styles.thin_border

            worksheet[f"C{row}"] = f"='Monthly Entry'!{col}{self.rows.savings_transfer}"
            worksheet[f"C{row}"].font = self.styles.calc_font()
            worksheet[f"C{row}"].border = self.styles.thin_border
            worksheet[f"C{row}"].alignment = self.styles.right_alignment()
            worksheet[f"C{row}"].number_format = self.config.currency_format

            worksheet[f"D{row}"] = f"='Monthly Entry'!{col}{self.rows.running_savings}"
            worksheet[f"D{row}"].font = self.styles.calc_font()
            worksheet[f"D{row}"].border = self.styles.thin_border
            worksheet[f"D{row}"].alignment = self.styles.right_alignment()
            worksheet[f"D{row}"].number_format = self.config.currency_format

            worksheet[f"E{row}"] = f"=IF($C$8=0,0,D{row}/$C$8)"
            worksheet[f"E{row}"].font = self.styles.calc_font()
            worksheet[f"E{row}"].border = self.styles.thin_border
            worksheet[f"E{row}"].alignment = self.styles.right_alignment()
            worksheet[f"E{row}"].number_format = self.config.percent_format

            if row % 2 == 0:
                for col_letter in ["B", "C", "D", "E"]:
                    worksheet[f"{col_letter}{row}"].fill = self.styles.alt_row_fill()
            row += 1

    def _create_guidelines(self, worksheet) -> None:
        row = 34
        worksheet[f"B{row}"] = "GUIDELINES"
        worksheet[f"B{row}"].font = Font(size=12, bold=True, color=self.styles.palette.header_dark)
        row += 2

        guidelines = [
            "• Emergency fund should cover 6 months of essential expenses",
            "• Keep emergency fund in a separate, easily accessible savings account",
            "• Only use for true emergencies (job loss, medical, major repairs)",
            "• Replenish immediately after any withdrawal",
        ]
        for guideline in guidelines:
            worksheet[f"B{row}"] = guideline
            worksheet[f"B{row}"].font = Font(size=10, color="555555")
            worksheet.merge_cells(f"B{row}:E{row}")
            row += 1
