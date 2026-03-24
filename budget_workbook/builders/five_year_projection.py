"""5-year projection worksheet builder."""

from openpyxl.styles import Font

from .base import BaseSheetBuilder


class FiveYearProjectionSheetBuilder(BaseSheetBuilder):
    """Build the 5-Year Projection worksheet."""

    def build(self, workbook):
        worksheet = workbook.create_sheet("5-Year Projection")
        worksheet.sheet_view.showGridLines = False
        for col, width in zip(["A", "B", "C", "D", "E", "F", "G", "H"], [3, 24, 18, 18, 18, 18, 18, 18]):
            worksheet.column_dimensions[col].width = width

        worksheet.merge_cells("B2:H2")
        worksheet["B2"] = "5-YEAR SAVINGS PROJECTION"
        worksheet["B2"].font = Font(size=18, bold=True, color=self.styles.palette.header_dark)
        worksheet["B2"].alignment = self.styles.left_alignment()
        worksheet.row_dimensions[2].height = 30

        worksheet.merge_cells("B3:H3")
        worksheet["B3"] = "Projection compares budget and actual savings trajectories using formula-driven assumptions."
        worksheet["B3"].font = Font(size=10, italic=True, color="666666")

        self._create_assumptions(worksheet)
        self._create_projection_table(worksheet)
        return worksheet

    def _create_assumptions(self, worksheet) -> None:
        worksheet["B5"] = "PROJECTION ASSUMPTIONS"
        worksheet["B5"].font = Font(size=12, bold=True, color=self.styles.palette.header_dark)
        rows = [
            ("Annual Growth Rate", self.config.annual_growth_rate, self.config.percent_format),
            ("Monthly Savings Average (Budget)", f"='Monthly Entry'!{self.year_budget_letter()}{self.rows.savings_transfer}/12", self.config.currency_format),
            ("Monthly Savings Average (Actual)", f"='Monthly Entry'!{self.year_actual_letter()}{self.rows.savings_transfer}/12", self.config.currency_format),
            ("Starting Balance (Budget)", f"='Monthly Entry'!{self.year_budget_letter()}{self.rows.running_savings}", self.config.currency_format),
            ("Starting Balance (Actual)", f"='Monthly Entry'!{self.year_actual_letter()}{self.rows.running_savings}", self.config.currency_format),
        ]
        row = 6
        for label, value, number_format in rows:
            worksheet[f"B{row}"] = label
            worksheet[f"B{row}"].font = Font(size=10, color=self.styles.palette.text_dark)
            worksheet[f"B{row}"].border = self.styles.thin_border
            worksheet[f"C{row}"] = value
            worksheet[f"C{row}"].font = self.styles.calc_font() if isinstance(value, str) and str(value).startswith("=") else self.styles.input_font()
            worksheet[f"C{row}"].border = self.styles.thin_border
            worksheet[f"C{row}"].alignment = self.styles.right_alignment()
            worksheet[f"C{row}"].number_format = number_format
            row += 1

    def _create_projection_table(self, worksheet) -> None:
        worksheet["B13"] = "YEAR-BY-YEAR PROJECTION"
        worksheet["B13"].font = Font(size=14, bold=True, color=self.styles.palette.header_dark)

        headers = ["Year", "Budget Start", "Actual Start", "Budget Contribution", "Actual Contribution", "Budget End", "Actual End", "Variance"]
        for index, header in enumerate(headers, start=2):
            cell = worksheet.cell(row=15, column=index)
            cell.value = header
            cell.font = self.styles.header_font()
            cell.fill = self.styles.header_fill()
            cell.border = self.styles.thin_border
            cell.alignment = self.styles.left_alignment() if index == 2 else self.styles.right_alignment()

        row = 16
        for year in range(1, 6):
            worksheet[f"B{row}"] = f"Year {year}"
            worksheet[f"B{row}"].font = Font(size=10, color=self.styles.palette.text_dark)
            worksheet[f"B{row}"].border = self.styles.thin_border
            worksheet[f"B{row}"].alignment = self.styles.left_alignment()
            if year == 1:
                worksheet[f"C{row}"] = "=C9"
                worksheet[f"D{row}"] = "=C10"
            else:
                worksheet[f"C{row}"] = f"=G{row-1}"
                worksheet[f"D{row}"] = f"=H{row-1}"
            worksheet[f"E{row}"] = "=C7*12"
            worksheet[f"F{row}"] = "=C8*12"
            worksheet[f"G{row}"] = f"=C{row}+E{row}+(C{row}+E{row}/2)*$C$6"
            worksheet[f"H{row}"] = f"=D{row}+F{row}+(D{row}+F{row}/2)*$C$6"
            worksheet[f"I{row}"] = f"=H{row}-G{row}"
            for col in ["C", "D", "E", "F", "G", "H", "I"]:
                worksheet[f"{col}{row}"].font = self.styles.calc_font()
                worksheet[f"{col}{row}"].border = self.styles.thin_border
                worksheet[f"{col}{row}"].alignment = self.styles.right_alignment()
                worksheet[f"{col}{row}"].number_format = self.config.currency_format
            row += 1
