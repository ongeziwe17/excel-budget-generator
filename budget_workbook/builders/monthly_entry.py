"""Monthly entry worksheet builder."""

from __future__ import annotations

from collections.abc import Iterable

from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

from .base import BaseSheetBuilder


class MonthlyEntrySheetBuilder(BaseSheetBuilder):
    """Build the Monthly Entry worksheet with budget vs actual tracking."""

    grocery_table_name = "GroceryDetailTable"
    custom_variable_table_name = "VariableExpenseTable"

    def build(self, workbook):
        worksheet = workbook.create_sheet("Monthly Entry")
        worksheet.sheet_view.showGridLines = False
        self._configure_columns(worksheet)
        self._create_title_block(worksheet)
        self.create_month_header_grid(worksheet, title_row=5, subtitle_row=6)
        self._create_core_sections(worksheet)
        grocery_table_end = self._create_grocery_detail_table(worksheet, start_row=49)
        custom_table_end = self._create_custom_variable_table(worksheet, start_row=grocery_table_end + 4)
        self._create_detail_notes(worksheet, custom_table_end + 3)
        self._apply_variance_formatting(worksheet, grocery_table_end, custom_table_end)
        return worksheet

    def _configure_columns(self, worksheet) -> None:
        worksheet.column_dimensions["A"].width = 3
        worksheet.column_dimensions["B"].width = 34
        for index in range(len(self.config.months)):
            worksheet.column_dimensions[get_column_letter(self.month_budget_col(index))].width = 12
            worksheet.column_dimensions[get_column_letter(self.month_actual_col(index))].width = 12
            worksheet.column_dimensions[get_column_letter(self.month_variance_col(index))].width = 12
        worksheet.column_dimensions[get_column_letter(self.year_budget_col())].width = 14
        worksheet.column_dimensions[get_column_letter(self.year_actual_col())].width = 14
        worksheet.column_dimensions[get_column_letter(self.year_variance_col())].width = 14

    def _create_title_block(self, worksheet) -> None:
        worksheet.merge_cells(start_row=2, start_column=2, end_row=2, end_column=self.last_data_col())
        worksheet["B2"] = "MONTHLY DATA ENTRY"
        worksheet["B2"].font = Font(size=18, bold=True, color=self.styles.palette.header_dark)
        worksheet["B2"].alignment = self.styles.left_alignment()
        worksheet.row_dimensions[2].height = 30

        worksheet.merge_cells(start_row=3, start_column=2, end_row=3, end_column=self.last_data_col())
        worksheet["B3"] = (
            "Enter planned amounts in Budget cells and real cash movement in Actual cells. Variance = Actual - Budget."
        )
        worksheet["B3"].font = Font(size=10, italic=True, color="666666")

    def _create_core_sections(self, worksheet) -> None:
        self.create_section_header(worksheet, 7, "INCOME")
        self.create_input_row(worksheet, self.rows.net_income, "Net Income Received")

        self.create_section_header(worksheet, 10, "FIXED EXPENSES")
        for row, label in [
            (self.rows.data, "Data"),
            (self.rows.transport, "Transport"),
            (self.rows.subscriptions, "Subscriptions"),
            (self.rows.tfg_debit, "TFG Debit"),
            (self.rows.family_support, "Family Support"),
        ]:
            self.create_input_row(worksheet, row, label, indent=True)

        self.create_formula_row(
            worksheet,
            self.rows.tithe,
            "Tithe (10% of Net Income)",
            lambda index, _month: f"={self.month_budget_letter(index)}{self.rows.net_income}*0.1",
            lambda index, _month: f"={self.month_actual_letter(index)}{self.rows.net_income}*0.1",
            indent=True,
        )
        self.create_input_row(worksheet, self.rows.rent_housing, "Rent / Housing", indent=True)

        self.create_section_header(worksheet, 19, "VARIABLE EXPENSES")
        self.create_formula_row(
            worksheet,
            self.rows.groceries,
            "Groceries (rolls up from detail table below)",
            lambda _index, month: f"=SUM({self.grocery_table_name}[{month} Budget])",
            lambda _index, month: f"=SUM({self.grocery_table_name}[{month} Actual])",
            indent=True,
        )
        for row, label in [
            (self.rows.eating_out, "Eating Out"),
            (self.rows.lunch, "Lunch"),
            (self.rows.haircuts, "Haircuts"),
            (self.rows.clothing, "Clothing"),
            (self.rows.random_spending, "Random Spending"),
        ]:
            self.create_input_row(worksheet, row, label, indent=True)
        self.create_formula_row(
            worksheet,
            self.rows.custom_variable_total,
            "Additional Variable Items (rolls up from detail table below)",
            lambda _index, month: f"=SUM({self.custom_variable_table_name}[{month} Budget])",
            lambda _index, month: f"=SUM({self.custom_variable_table_name}[{month} Actual])",
            indent=True,
        )

        self.create_section_header(worksheet, 28, "SAVINGS & TRANSFERS")
        self.create_input_row(worksheet, self.rows.savings_transfer, "Transfer to Savings")
        self.create_input_row(worksheet, self.rows.starting_savings, "Starting Savings Balance")

        self.create_section_header(worksheet, 32, "BONUSES RECEIVED")
        self.create_input_row(worksheet, self.rows.bonus, "Bonus Received")

        self.create_section_header(worksheet, 35, "UNEXPECTED EXPENSES")
        self.create_input_row(worksheet, self.rows.unexpected_expenses, "Unexpected Expenses")

        self.create_section_header(worksheet, 38, "SUMMARY TOTALS")
        self._create_total_rows(worksheet)

    def _create_total_rows(self, worksheet) -> None:
        self.create_formula_row(
            worksheet,
            self.rows.total_fixed,
            "Total Fixed Expenses",
            lambda index, _month: self._sum_row_formula(
                index,
                [
                    self.rows.data,
                    self.rows.transport,
                    self.rows.subscriptions,
                    self.rows.tfg_debit,
                    self.rows.family_support,
                    self.rows.tithe,
                    self.rows.rent_housing,
                ],
                actual=False,
            ),
            lambda index, _month: self._sum_row_formula(
                index,
                [
                    self.rows.data,
                    self.rows.transport,
                    self.rows.subscriptions,
                    self.rows.tfg_debit,
                    self.rows.family_support,
                    self.rows.tithe,
                    self.rows.rent_housing,
                ],
                actual=True,
            ),
            fill=self.styles.section_fill(),
            font_color=self.styles.palette.white,
            bold=True,
        )
        self.create_formula_row(
            worksheet,
            self.rows.total_variable,
            "Total Variable Expenses",
            lambda index, _month: self._sum_row_formula(
                index,
                [
                    self.rows.groceries,
                    self.rows.eating_out,
                    self.rows.lunch,
                    self.rows.haircuts,
                    self.rows.clothing,
                    self.rows.random_spending,
                    self.rows.custom_variable_total,
                ],
                actual=False,
            ),
            lambda index, _month: self._sum_row_formula(
                index,
                [
                    self.rows.groceries,
                    self.rows.eating_out,
                    self.rows.lunch,
                    self.rows.haircuts,
                    self.rows.clothing,
                    self.rows.random_spending,
                    self.rows.custom_variable_total,
                ],
                actual=True,
            ),
            fill=self.styles.section_fill(),
            font_color=self.styles.palette.white,
            bold=True,
        )
        self.create_formula_row(
            worksheet,
            self.rows.total_expenses,
            "TOTAL ALL EXPENSES",
            lambda index, _month: self._sum_row_formula(
                index,
                [self.rows.total_fixed, self.rows.total_variable, self.rows.unexpected_expenses],
                actual=False,
            ),
            lambda index, _month: self._sum_row_formula(
                index,
                [self.rows.total_fixed, self.rows.total_variable, self.rows.unexpected_expenses],
                actual=True,
            ),
            fill=self.styles.header_fill(),
            font_color=self.styles.palette.white,
            bold=True,
        )
        self.create_formula_row(
            worksheet,
            self.rows.total_income,
            "TOTAL INCOME",
            lambda index, _month: self._sum_row_formula(index, [self.rows.net_income, self.rows.bonus], actual=False),
            lambda index, _month: self._sum_row_formula(index, [self.rows.net_income, self.rows.bonus], actual=True),
            fill=self.styles.positive_fill(),
            font_color=self.styles.palette.white,
            bold=True,
        )
        self.create_formula_row(
            worksheet,
            self.rows.surplus_deficit,
            "SURPLUS / (DEFICIT)",
            lambda index, _month: f"={self.month_budget_letter(index)}{self.rows.total_income}-{self.month_budget_letter(index)}{self.rows.total_expenses}-{self.month_budget_letter(index)}{self.rows.savings_transfer}",
            lambda index, _month: f"={self.month_actual_letter(index)}{self.rows.total_income}-{self.month_actual_letter(index)}{self.rows.total_expenses}-{self.month_actual_letter(index)}{self.rows.savings_transfer}",
            fill=self.styles.section_fill(),
            font_color=self.styles.palette.white,
            bold=True,
        )
        self.create_formula_row(
            worksheet,
            self.rows.savings_rate,
            "Savings Rate %",
            lambda index, _month: f"=IF({self.month_budget_letter(index)}{self.rows.total_income}=0,0,{self.month_budget_letter(index)}{self.rows.savings_transfer}/{self.month_budget_letter(index)}{self.rows.total_income})",
            lambda index, _month: f"=IF({self.month_actual_letter(index)}{self.rows.total_income}=0,0,{self.month_actual_letter(index)}{self.rows.savings_transfer}/{self.month_actual_letter(index)}{self.rows.total_income})",
            use_percent=True,
        )
        self._create_running_savings_row(worksheet)

    def _create_running_savings_row(self, worksheet) -> None:
        self._set_row_label(worksheet, self.rows.running_savings, "Running Savings Balance", bold=True)
        for index, _month in enumerate(self.config.months):
            budget_col = self.month_budget_col(index)
            actual_col = self.month_actual_col(index)
            variance_col = self.month_variance_col(index)
            if index == 0:
                budget_formula = f"={self.month_budget_letter(index)}{self.rows.starting_savings}+{self.month_budget_letter(index)}{self.rows.savings_transfer}"
                actual_formula = f"={self.month_actual_letter(index)}{self.rows.starting_savings}+{self.month_actual_letter(index)}{self.rows.savings_transfer}"
            else:
                previous_budget_col = get_column_letter(self.month_budget_col(index - 1))
                previous_actual_col = get_column_letter(self.month_actual_col(index - 1))
                budget_formula = f"={previous_budget_col}{self.rows.running_savings}+{self.month_budget_letter(index)}{self.rows.savings_transfer}"
                actual_formula = f"={previous_actual_col}{self.rows.running_savings}+{self.month_actual_letter(index)}{self.rows.savings_transfer}"

            worksheet.cell(row=self.rows.running_savings, column=budget_col).value = budget_formula
            worksheet.cell(row=self.rows.running_savings, column=actual_col).value = actual_formula
            worksheet.cell(row=self.rows.running_savings, column=variance_col).value = f"={get_column_letter(actual_col)}{self.rows.running_savings}-{get_column_letter(budget_col)}{self.rows.running_savings}"
            for col in (budget_col, actual_col, variance_col):
                cell = worksheet.cell(row=self.rows.running_savings, column=col)
                cell.font = self.styles.calc_font()
                cell.border = self.styles.thin_border
                cell.alignment = self.styles.right_alignment()
                cell.number_format = self.config.currency_format
        worksheet.cell(row=self.rows.running_savings, column=self.year_budget_col()).value = f"={self.month_budget_letter(len(self.config.months)-1)}{self.rows.running_savings}"
        worksheet.cell(row=self.rows.running_savings, column=self.year_actual_col()).value = f"={self.month_actual_letter(len(self.config.months)-1)}{self.rows.running_savings}"
        worksheet.cell(row=self.rows.running_savings, column=self.year_variance_col()).value = f"={self.year_actual_letter()}{self.rows.running_savings}-{self.year_budget_letter()}{self.rows.running_savings}"
        for col in (self.year_budget_col(), self.year_actual_col(), self.year_variance_col()):
            cell = worksheet.cell(row=self.rows.running_savings, column=col)
            cell.font = self.styles.calc_font()
            cell.border = self.styles.thin_border
            cell.alignment = self.styles.right_alignment()
            cell.number_format = self.config.currency_format

    def _create_grocery_detail_table(self, worksheet, start_row: int) -> int:
        self.create_section_header(worksheet, 47, "DETAIL TABLES")
        self.create_section_header(worksheet, 48, "GROCERY DETAIL (ROLLS UP TO GROCERIES ROW ABOVE)")
        headers = ["Group", "Item", *self._detail_value_headers(), "Year Budget", "Year Actual", "Year Variance"]
        for index, header in enumerate(headers, start=2):
            cell = worksheet.cell(row=start_row, column=index)
            cell.value = header
            cell.font = self.styles.header_font()
            cell.fill = self.styles.header_fill()
            cell.border = self.styles.thin_border
            cell.alignment = self.styles.center_alignment() if index > 3 else self.styles.left_alignment()

        current_row = start_row + 1
        for group, items in self._grocery_seed_data():
            for item in items:
                self._write_detail_row(worksheet, current_row, group, item)
                current_row += 1
            for _ in range(3):
                self._write_detail_row(worksheet, current_row, group, "")
                current_row += 1

        self._add_table(
            worksheet,
            self.grocery_table_name,
            start_row,
            current_row - 1,
            len(headers) + 1,
        )
        return current_row - 1

    def _create_custom_variable_table(self, worksheet, start_row: int) -> int:
        self.create_section_header(worksheet, start_row - 1, "ADDITIONAL VARIABLE ITEMS (ROLLS UP TO CUSTOM VARIABLE TOTAL)")
        headers = ["Category", "Item", *self._detail_value_headers(), "Year Budget", "Year Actual", "Year Variance"]
        for index, header in enumerate(headers, start=2):
            cell = worksheet.cell(row=start_row, column=index)
            cell.value = header
            cell.font = self.styles.header_font()
            cell.fill = self.styles.header_fill()
            cell.border = self.styles.thin_border
            cell.alignment = self.styles.center_alignment() if index > 3 else self.styles.left_alignment()

        current_row = start_row + 1
        for category in ["Personal Care", "Entertainment", "Gifts", "Medical", "Other"]:
            self._write_detail_row(worksheet, current_row, category, "")
            current_row += 1

        self._add_table(
            worksheet,
            self.custom_variable_table_name,
            start_row,
            current_row - 1,
            len(headers) + 1,
        )
        return current_row - 1

    def _create_detail_notes(self, worksheet, start_row: int) -> None:
        worksheet[f"B{start_row}"] = "DETAIL SECTION NOTES"
        worksheet[f"B{start_row}"].font = Font(size=12, bold=True, color=self.styles.palette.header_dark)
        notes = [
            "• Enter Budget and Actual values only in the white table cells; Variance and yearly totals calculate automatically.",
            "• Add new grocery rows inside GroceryDetailTable to keep Groceries totals, dashboards, and charts in sync.",
            "• Add new custom variable rows inside VariableExpenseTable to extend the model without changing formulas.",
            "• To support more months, extend WorkbookConfig.months and regenerate the workbook so every sheet and chart expands safely.",
        ]
        for offset, note in enumerate(notes, start=2):
            worksheet[f"B{start_row + offset}"] = note
            worksheet[f"B{start_row + offset}"].font = Font(size=10, color="555555")
            worksheet.merge_cells(start_row=start_row + offset, start_column=2, end_row=start_row + offset, end_column=self.last_data_col())

    def _write_detail_row(self, worksheet, row: int, group: str, item: str) -> None:
        worksheet[f"B{row}"] = group
        worksheet[f"C{row}"] = item
        worksheet[f"B{row}"].font = Font(size=10, color=self.styles.palette.text_dark)
        worksheet[f"C{row}"].font = Font(size=10, color=self.styles.palette.text_dark)
        worksheet[f"B{row}"].border = self.styles.thin_border
        worksheet[f"C{row}"].border = self.styles.thin_border
        worksheet[f"B{row}"].alignment = self.styles.left_alignment()
        worksheet[f"C{row}"].alignment = self.styles.left_alignment()

        for index, month in enumerate(self.config.months):
            budget_col = self.month_budget_col(index) + 1
            actual_col = self.month_actual_col(index) + 1
            variance_col = self.month_variance_col(index) + 1
            budget_letter = get_column_letter(budget_col)
            actual_letter = get_column_letter(actual_col)
            variance_letter = get_column_letter(variance_col)
            worksheet[f"{budget_letter}{row}"].font = self.styles.input_font()
            worksheet[f"{actual_letter}{row}"].font = self.styles.input_font()
            worksheet[f"{variance_letter}{row}"].font = self.styles.calc_font()
            worksheet[f"{variance_letter}{row}"] = f"={actual_letter}{row}-{budget_letter}{row}"
            for letter in (budget_letter, actual_letter, variance_letter):
                worksheet[f"{letter}{row}"].border = self.styles.thin_border
                worksheet[f"{letter}{row}"].alignment = self.styles.right_alignment()
                worksheet[f"{letter}{row}"].number_format = self.config.currency_format

        year_budget_letter = get_column_letter(self.year_budget_col() + 1)
        year_actual_letter = get_column_letter(self.year_actual_col() + 1)
        year_variance_letter = get_column_letter(self.year_variance_col() + 1)
        budget_cells = [f"{get_column_letter(self.month_budget_col(index) + 1)}{row}" for index in range(len(self.config.months))]
        actual_cells = [f"{get_column_letter(self.month_actual_col(index) + 1)}{row}" for index in range(len(self.config.months))]
        worksheet[f"{year_budget_letter}{row}"] = f"=SUM({','.join(budget_cells)})"
        worksheet[f"{year_actual_letter}{row}"] = f"=SUM({','.join(actual_cells)})"
        worksheet[f"{year_variance_letter}{row}"] = f"={year_actual_letter}{row}-{year_budget_letter}{row}"
        for letter in (year_budget_letter, year_actual_letter, year_variance_letter):
            worksheet[f"{letter}{row}"].font = self.styles.calc_font()
            worksheet[f"{letter}{row}"].border = self.styles.thin_border
            worksheet[f"{letter}{row}"].alignment = self.styles.right_alignment()
            worksheet[f"{letter}{row}"].number_format = self.config.currency_format

    def _detail_value_headers(self) -> Iterable[str]:
        for month in self.config.months:
            yield f"{month} Budget"
            yield f"{month} Actual"
            yield f"{month} Variance"

    def _add_table(self, worksheet, table_name: str, start_row: int, end_row: int, end_col: int) -> None:
        table = Table(
            displayName=table_name,
            ref=f"B{start_row}:{get_column_letter(end_col)}{end_row}",
        )
        table.tableStyleInfo = TableStyleInfo(
            name="TableStyleMedium2",
            showFirstColumn=False,
            showLastColumn=False,
            showRowStripes=True,
            showColumnStripes=False,
        )
        worksheet.add_table(table)

    def _apply_variance_formatting(self, worksheet, grocery_table_end: int, custom_table_end: int) -> None:
        self.add_variance_conditional_formatting(worksheet, self.rows.data, self.rows.total_expenses)
        self.add_variance_conditional_formatting(worksheet, self.rows.groceries, self.rows.custom_variable_total)
        detail_variance_columns = [self.month_variance_col(index) + 1 for index in range(len(self.config.months))]
        detail_variance_columns.append(self.year_variance_col() + 1)
        self.add_variance_conditional_formatting(worksheet, 50, grocery_table_end, columns=detail_variance_columns)
        self.add_variance_conditional_formatting(worksheet, grocery_table_end + 5, custom_table_end, columns=detail_variance_columns)

    def _sum_row_formula(self, index: int, rows: list[int], *, actual: bool) -> str:
        col_letter = self.month_actual_letter(index) if actual else self.month_budget_letter(index)
        return "=" + "+".join(f"{col_letter}{row}" for row in rows)

    def _grocery_seed_data(self) -> list[tuple[str, list[str]]]:
        return [
            (
                "Grains / Staples",
                [
                    "Rice (2KG x 2)",
                    "Maize Meal (2.5KG)",
                    "Mr Pasta (500G x 2)",
                    "Spaghetti",
                    "Spar Noodles",
                    "Samp and Beans (500G)",
                ],
            ),
            (
                "Protein / Meat",
                [
                    "Farmers Choice Chicken Braai Pack (+-2000g)",
                    "Beef (1kg)",
                    "Spar Mince",
                    "Rise Spar Sausage",
                    "Fish (410G x 2)",
                    "Baked Beans (410G x 2)",
                    "Eggs (30)",
                    "Peanut Butter YUM YUM",
                ],
            ),
            (
                "Vegetables",
                [
                    "McCain Frozen Mixed Veg (250g)",
                    "McCain Frozen Peas (250g)",
                    "McCain Garden Mix (250g)",
                    "SweetCorn (4)",
                    "Potatoes (4kg)",
                    "Carrots (3kg)",
                    "Butternut (4kg)",
                    "Onion (1kg)",
                    "Ginger Garlic (150g)",
                    "Mixed Pepper (4)",
                    "Cucumber",
                ],
            ),
            (
                "Spices / Condiments",
                [
                    "Steak and Chops Spice (160g)",
                    "Chicken Spice (160g)",
                    "BBQ Spice (160g)",
                    "Black Pepper (160g)",
                    "Aromat (200g)",
                    "6 Gun Spice (200g)",
                    "Rajah Mild 'n' Spicy (160g)",
                    "Mayonnaise Nola (750g)",
                    "Wellington Tomato Sauce",
                    "Wellington Sweet Chilli (500ml)",
                    "Tomato Paste (2 x 50g)",
                    "Champion Braai (375ml)",
                    "Chutney (470g)",
                ],
            ),
            (
                "Cereals / Breakfast",
                [
                    "Corn Flakes (1kg)",
                    "Coco Pops (350g)",
                    "Nutrific (200g)",
                    "Muesli (350g)",
                    "Future Life (500g)",
                ],
            ),
            (
                "Fruit",
                [
                    "Apples (2kg)",
                    "Bananas (1.2kg)",
                    "Spar Peaches",
                    "Frozen Berries (350g)",
                    "Canned Fruit (850g)",
                    "Lemon (800g)",
                ],
            ),
            (
                "Dairy / Fridge Items",
                [
                    "Milk (1L x 6)",
                    "Cheese",
                    "Rama",
                ],
            ),
            (
                "Household / Cooking Items",
                [
                    "Sugar Selati (2kg)",
                ],
            ),
            (
                "Snacks",
                [
                    "Ice Cream (80ml)",
                    "Doritos (2 x 145g)",
                ],
            ),
        ]
