"""Trend analysis worksheet builder."""

from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from .base import BaseSheetBuilder


class TrendAnalysisSheetBuilder(BaseSheetBuilder):
    """Build the Trend Analysis worksheet."""

    def build(self, workbook):
        worksheet = workbook.create_sheet("Trend Analysis")
        worksheet.sheet_view.showGridLines = False
        self._configure_columns(worksheet)
        self._create_header(worksheet)
        month_section_end = self._create_monthly_variance_section(worksheet, start_row=6)
        self._create_change_section(worksheet, start_row=month_section_end + 3)
        return worksheet

    def _configure_columns(self, worksheet) -> None:
        worksheet.column_dimensions["A"].width = 3
        worksheet.column_dimensions["B"].width = 30
        for index in range(len(self.config.months)):
            worksheet.column_dimensions[get_column_letter(3 + index)].width = 14
        for col in ["P", "Q", "R", "S", "T"]:
            worksheet.column_dimensions[col].width = 14

    def _create_header(self, worksheet) -> None:
        last_col = get_column_letter(2 + len(self.config.months))
        worksheet.merge_cells(f"B2:{last_col}2")
        worksheet["B2"] = "TREND ANALYSIS"
        worksheet["B2"].font = Font(size=18, bold=True, color=self.styles.palette.header_dark)
        worksheet["B2"].alignment = self.styles.left_alignment()
        worksheet.row_dimensions[2].height = 30

        worksheet.merge_cells(f"B3:{last_col}3")
        worksheet["B3"] = "Budget vs Actual variance and month-over-month movement driven from Monthly Entry totals."
        worksheet["B3"].font = Font(size=10, italic=True, color="666666")

    def _create_monthly_variance_section(self, worksheet, start_row: int) -> int:
        worksheet[f"B{start_row}"] = "MONTHLY BUDGET VS ACTUAL"
        worksheet[f"B{start_row}"].font = Font(size=14, bold=True, color=self.styles.palette.header_dark)
        header_row = start_row + 2
        worksheet[f"B{header_row}"] = "Metric"
        worksheet[f"B{header_row}"].font = self.styles.header_font()
        worksheet[f"B{header_row}"].fill = self.styles.header_fill()
        worksheet[f"B{header_row}"].border = self.styles.thin_border
        worksheet[f"B{header_row}"].alignment = self.styles.left_alignment()
        for index, month in enumerate(self.config.months):
            cell = worksheet.cell(row=header_row, column=3 + index)
            cell.value = month
            cell.font = self.styles.header_font()
            cell.fill = self.styles.header_fill()
            cell.border = self.styles.thin_border
            cell.alignment = self.styles.center_alignment()

        metrics = [
            ("Total Expense Variance", self.rows.total_expenses),
            ("Groceries Variance", self.rows.groceries),
            ("Additional Variable Items Variance", self.rows.custom_variable_total),
            ("Unexpected Expense Variance", self.rows.unexpected_expenses),
            ("Total Income Variance", self.rows.total_income),
            ("Bonus Variance", self.rows.bonus),
            ("Savings Transfer Variance", self.rows.savings_transfer),
            ("Savings Rate Variance", self.rows.savings_rate),
        ]
        row = header_row + 1
        for label, source_row in metrics:
            worksheet[f"B{row}"] = label
            worksheet[f"B{row}"].font = Font(size=10, color=self.styles.palette.text_dark)
            worksheet[f"B{row}"].border = self.styles.thin_border
            worksheet[f"B{row}"].alignment = self.styles.left_alignment()
            for index in range(len(self.config.months)):
                cell = worksheet.cell(row=row, column=3 + index)
                cell.value = f"='Monthly Entry'!{self.month_variance_letter(index)}{source_row}"
                cell.font = self.styles.calc_font()
                cell.border = self.styles.thin_border
                cell.alignment = self.styles.right_alignment()
                cell.number_format = self.config.percent_format if source_row == self.rows.savings_rate else self.config.currency_format
            if row % 2 == 0:
                worksheet[f"B{row}"].fill = self.styles.alt_row_fill()
            row += 1

        red_fill = PatternFill(start_color=self.styles.palette.red_fill, end_color=self.styles.palette.red_fill, fill_type="solid")
        green_fill = PatternFill(start_color=self.styles.palette.green_fill, end_color=self.styles.palette.green_fill, fill_type="solid")
        for col in range(3, 3 + len(self.config.months)):
            cell_range = f"{get_column_letter(col)}{header_row + 1}:{get_column_letter(col)}{row - 1}"
            worksheet.conditional_formatting.add(cell_range, CellIsRule(operator="greaterThan", formula=["0"], fill=red_fill))
            worksheet.conditional_formatting.add(cell_range, CellIsRule(operator="lessThan", formula=["0"], fill=green_fill))
        return row - 1

    def _create_change_section(self, worksheet, start_row: int) -> None:
        worksheet[f"B{start_row}"] = "MONTH-OVER-MONTH ACTUAL CHANGE"
        worksheet[f"B{start_row}"].font = Font(size=14, bold=True, color=self.styles.palette.header_dark)
        header_row = start_row + 2
        worksheet[f"B{header_row}"] = "Metric"
        worksheet[f"B{header_row}"].font = self.styles.header_font()
        worksheet[f"B{header_row}"].fill = self.styles.header_fill()
        worksheet[f"B{header_row}"].border = self.styles.thin_border
        worksheet[f"B{header_row}"].alignment = self.styles.left_alignment()

        change_labels = [f"{self.config.months[index]}-{self.config.months[index-1]}" for index in range(1, len(self.config.months))]
        for index, label in enumerate(change_labels):
            cell = worksheet.cell(row=header_row, column=3 + index)
            cell.value = label
            cell.font = self.styles.header_font()
            cell.fill = self.styles.header_fill()
            cell.border = self.styles.thin_border
            cell.alignment = self.styles.center_alignment()

        metrics = [
            ("Actual Total Expenses Change", self.rows.total_expenses, self.config.currency_format),
            ("Actual Total Income Change", self.rows.total_income, self.config.currency_format),
            ("Actual Savings Transfer Change", self.rows.savings_transfer, self.config.currency_format),
            ("Actual Groceries Change", self.rows.groceries, self.config.currency_format),
            ("Actual Savings Rate Change", self.rows.savings_rate, self.config.percent_format),
        ]
        row = header_row + 1
        for label, source_row, number_format in metrics:
            worksheet[f"B{row}"] = label
            worksheet[f"B{row}"].font = Font(size=10, color=self.styles.palette.text_dark)
            worksheet[f"B{row}"].border = self.styles.thin_border
            worksheet[f"B{row}"].alignment = self.styles.left_alignment()
            for index in range(1, len(self.config.months)):
                current_letter = self.month_actual_letter(index)
                previous_letter = self.month_actual_letter(index - 1)
                cell = worksheet.cell(row=row, column=2 + index)
                cell.value = f"='Monthly Entry'!{current_letter}{source_row}-'Monthly Entry'!{previous_letter}{source_row}"
                cell.font = self.styles.calc_font()
                cell.border = self.styles.thin_border
                cell.alignment = self.styles.right_alignment()
                cell.number_format = number_format
            row += 1

        red_fill = PatternFill(start_color=self.styles.palette.red_fill, end_color=self.styles.palette.red_fill, fill_type="solid")
        green_fill = PatternFill(start_color=self.styles.palette.green_fill, end_color=self.styles.palette.green_fill, fill_type="solid")
        for col in range(3, 2 + len(self.config.months)):
            cell_range = f"{get_column_letter(col)}{header_row + 1}:{get_column_letter(col)}{row - 1}"
            worksheet.conditional_formatting.add(cell_range, CellIsRule(operator="greaterThan", formula=["0"], fill=red_fill))
            worksheet.conditional_formatting.add(cell_range, CellIsRule(operator="lessThan", formula=["0"], fill=green_fill))
