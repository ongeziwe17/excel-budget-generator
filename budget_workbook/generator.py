"""Application service for budget workbook generation."""

from pathlib import Path

from openpyxl import Workbook

from budget_workbook.builders.bonus_tracker import BonusTrackerSheetBuilder
from budget_workbook.builders.couple import (
    HouseholdBonusTrackerSheetBuilder,
    HouseholdCoverSheetBuilder,
    HouseholdDashboardSheetBuilder,
    HouseholdProjectionSheetBuilder,
    HouseholdSavingsSheetBuilder,
    HouseholdSetupSheetBuilder,
    HouseholdTrendSheetBuilder,
    PartnerBudgetSheetBuilder,
    SharedHouseholdSheetBuilder,
)
from budget_workbook.builders.cover import CoverSheetBuilder
from budget_workbook.builders.emergency_fund import EmergencyFundSheetBuilder
from budget_workbook.builders.five_year_projection import FiveYearProjectionSheetBuilder
from budget_workbook.builders.monthly_entry import MonthlyEntrySheetBuilder
from budget_workbook.builders.summary_dashboard import SummaryDashboardSheetBuilder
from budget_workbook.builders.trend_analysis import TrendAnalysisSheetBuilder
from budget_workbook.config import WorkbookConfig, WorkbookMode
from budget_workbook.rows import MonthlyEntryRows
from budget_workbook.styles import WorkbookStyles
from budget_workbook.versioning import next_available_patch_version, parse_versioned_filename


class BudgetWorkbookGenerator:
    """Create the complete personal budget workbook."""

    def __init__(self, config: WorkbookConfig | None = None) -> None:
        self.config = config or WorkbookConfig()
        self.styles = WorkbookStyles()
        self.rows = MonthlyEntryRows()
        self.sheet_builders = self._create_sheet_builders()

    def _create_sheet_builders(self):
        if self.config.mode is WorkbookMode.COUPLE:
            return [
                HouseholdCoverSheetBuilder(self.config, self.styles, self.rows),
                HouseholdDashboardSheetBuilder(self.config, self.styles, self.rows),
                HouseholdSetupSheetBuilder(self.config, self.styles, self.rows),
                PartnerBudgetSheetBuilder(
                    self.config,
                    self.styles,
                    self.rows,
                    self.config.couple.partner_one,
                    1,
                ),
                PartnerBudgetSheetBuilder(
                    self.config,
                    self.styles,
                    self.rows,
                    self.config.couple.partner_two,
                    2,
                ),
                SharedHouseholdSheetBuilder(self.config, self.styles, self.rows),
                HouseholdTrendSheetBuilder(self.config, self.styles, self.rows),
                HouseholdSavingsSheetBuilder(self.config, self.styles, self.rows),
                HouseholdBonusTrackerSheetBuilder(self.config, self.styles, self.rows),
                HouseholdProjectionSheetBuilder(self.config, self.styles, self.rows),
            ]
        return [
            CoverSheetBuilder(self.config, self.styles, self.rows),
            MonthlyEntrySheetBuilder(self.config, self.styles, self.rows),
            SummaryDashboardSheetBuilder(self.config, self.styles, self.rows),
            TrendAnalysisSheetBuilder(self.config, self.styles, self.rows),
            BonusTrackerSheetBuilder(self.config, self.styles, self.rows),
            EmergencyFundSheetBuilder(self.config, self.styles, self.rows),
            FiveYearProjectionSheetBuilder(self.config, self.styles, self.rows),
        ]

    def create_workbook(self, output_path: str | Path | None = None) -> Path:
        """Generate the workbook and return the saved file path.

        Default generation auto-increments patch versions when prior workbooks exist.
        Custom non-versioned paths still protect against overwriting existing files.
        """
        target_path = self._resolve_output_path(output_path)
        target_path.parent.mkdir(parents=True, exist_ok=True)

        workbook = Workbook()
        for builder in self.sheet_builders:
            builder.build(workbook)

        workbook.calculation.calcMode = "auto"
        workbook.calculation.fullCalcOnLoad = True
        workbook.calculation.forceFullCalc = True

        workbook.save(target_path)
        return target_path

    def _resolve_output_path(self, output_path: str | Path | None) -> Path:
        if output_path is None:
            version = next_available_patch_version(
                directory=self.config.output_dir,
                prefix=self.config.workbook_name_prefix,
                start_version=self.config.workbook_version,
            )
            return self.config.output_dir / f"{self.config.workbook_name_prefix}_{version}.xlsx"

        target = Path(output_path)
        target.parent.mkdir(parents=True, exist_ok=True)

        parsed = parse_versioned_filename(target.name, self.config.workbook_name_prefix)
        if parsed is not None:
            version = next_available_patch_version(
                directory=target.parent,
                prefix=self.config.workbook_name_prefix,
                start_version=parsed,
            )
            return target.parent / f"{self.config.workbook_name_prefix}_{version}.xlsx"

        if target.exists():
            raise FileExistsError(
                f"Workbook already exists at {target}. "
                "Provide a new custom output path or use versioned naming."
            )
        return target
