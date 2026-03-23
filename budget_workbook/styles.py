"""Workbook styling primitives and helpers."""

from dataclasses import dataclass

from openpyxl.styles import Alignment, Border, Font, PatternFill, Side


@dataclass(frozen=True)
class WorkbookPalette:
    """Color palette shared by every worksheet."""

    header_dark: str = "1F4E79"
    header_light: str = "2E75B6"
    accent_blue: str = "5B9BD5"
    bg_light: str = "F5F5F5"
    text_dark: str = "000000"
    white: str = "FFFFFF"
    positive_green: str = "27AE60"
    negative_red: str = "E74C3C"
    input_blue: str = "0066CC"
    green_fill: str = "D4EDDA"
    red_fill: str = "F8D7DA"


class WorkbookStyles:
    """Factory for reusable style objects."""

    def __init__(self, palette: WorkbookPalette | None = None) -> None:
        self.palette = palette or WorkbookPalette()
        self.thin_border = Border(
            left=Side(style="thin", color="D0D0D0"),
            right=Side(style="thin", color="D0D0D0"),
            top=Side(style="thin", color="D0D0D0"),
            bottom=Side(style="thin", color="D0D0D0"),
        )

    def header_font(self) -> Font:
        return Font(color=self.palette.white, bold=True, size=11)

    def section_font(self) -> Font:
        return Font(color=self.palette.white, bold=True, size=10)

    def input_font(self) -> Font:
        return Font(color=self.palette.input_blue, size=11)

    def calc_font(self) -> Font:
        return Font(color=self.palette.text_dark, size=11)

    def header_fill(self) -> PatternFill:
        return PatternFill(
            start_color=self.palette.header_dark,
            end_color=self.palette.header_dark,
            fill_type="solid",
        )

    def section_fill(self) -> PatternFill:
        return PatternFill(
            start_color=self.palette.header_light,
            end_color=self.palette.header_light,
            fill_type="solid",
        )

    def alt_row_fill(self) -> PatternFill:
        return PatternFill(
            start_color=self.palette.bg_light,
            end_color=self.palette.bg_light,
            fill_type="solid",
        )

    def positive_fill(self) -> PatternFill:
        return PatternFill(
            start_color=self.palette.positive_green,
            end_color=self.palette.positive_green,
            fill_type="solid",
        )

    def negative_fill(self) -> PatternFill:
        return PatternFill(
            start_color=self.palette.negative_red,
            end_color=self.palette.negative_red,
            fill_type="solid",
        )

    @staticmethod
    def left_alignment() -> Alignment:
        return Alignment(horizontal="left", vertical="center")

    @staticmethod
    def right_alignment() -> Alignment:
        return Alignment(horizontal="right", vertical="center")

    @staticmethod
    def center_alignment() -> Alignment:
        return Alignment(horizontal="center", vertical="center")
