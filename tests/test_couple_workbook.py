"""Tests for couple and household workbook generation."""

from dataclasses import replace
from pathlib import Path
import re
from tempfile import TemporaryDirectory
import unittest

from openpyxl import load_workbook

from budget_workbook import (
    BudgetWorkbookGenerator,
    ContributionMethod,
    CoupleConfig,
    PersonConfig,
    WorkbookConfig,
    WorkbookMode,
)
from budget_workbook.builders.couple import PARTNER_ROWS, SHARED_ROWS


class CoupleConfigurationTests(unittest.TestCase):
    def test_net_income_cannot_exceed_gross_income(self) -> None:
        with self.assertRaisesRegex(ValueError, "cannot exceed"):
            PersonConfig(
                name="Alex",
                monthly_gross_income=20_000,
                monthly_net_income=21_000,
            )

    def test_partner_names_must_be_unique(self) -> None:
        with self.assertRaisesRegex(ValueError, "must be different"):
            CoupleConfig(
                partner_one=PersonConfig(name="Alex"),
                partner_two=PersonConfig(name="alex"),
            )

    def test_custom_contribution_shares_must_add_to_one(self) -> None:
        with self.assertRaisesRegex(ValueError, "add up to 100%"):
            CoupleConfig(
                partner_one=PersonConfig(name="One", custom_contribution_share=0.6),
                partner_two=PersonConfig(name="Two", custom_contribution_share=0.5),
                contribution_method=ContributionMethod.CUSTOM,
            )

    def test_valid_custom_contribution_shares(self) -> None:
        couple = CoupleConfig(
            partner_one=PersonConfig(name="One", custom_contribution_share=0.6),
            partner_two=PersonConfig(name="Two", custom_contribution_share=0.4),
            contribution_method=ContributionMethod.CUSTOM,
        )
        self.assertEqual(couple.partner_one.custom_contribution_share, 0.6)


class CoupleWorkbookTests(unittest.TestCase):
    expected_couple_sheets = [
        "Cover",
        "Household Dashboard",
        "Household Setup",
        "Partner 1 Budget",
        "Partner 2 Budget",
        "Shared Household",
        "Trend Analysis",
        "Savings Goals",
        "Bonus Tracker",
        "5-Year Projection",
    ]

    def setUp(self) -> None:
        self.temp_dir = TemporaryDirectory()
        self.addCleanup(self.temp_dir.cleanup)
        self.root = Path(self.temp_dir.name)
        couple = CoupleConfig(
            partner_one=PersonConfig(
                name="Alex",
                monthly_gross_income=35_000,
                monthly_net_income=30_000,
            ),
            partner_two=PersonConfig(name="Sam", monthly_net_income=20_000),
        )
        self.config = WorkbookConfig(
            mode=WorkbookMode.COUPLE,
            couple=couple,
            workbook_name_prefix="Household_Budget_Workbook",
            output_dir=self.root,
        )

    def generate(self, config: WorkbookConfig | None = None):
        path = self.root / "test-household.xlsx"
        BudgetWorkbookGenerator(config or self.config).create_workbook(path)
        return load_workbook(path, data_only=False)

    def test_couple_mode_creates_expected_sheet_order(self) -> None:
        workbook = self.generate()
        self.assertEqual(workbook.sheetnames, self.expected_couple_sheets)

    def test_single_mode_remains_backward_compatible(self) -> None:
        config = WorkbookConfig(output_dir=self.root)
        path = self.root / "test-personal.xlsx"
        BudgetWorkbookGenerator(config).create_workbook(path)
        workbook = load_workbook(path, data_only=False)
        self.assertEqual(
            workbook.sheetnames,
            [
                "Cover",
                "Monthly Entry",
                "Summary Dashboard",
                "Trend Analysis",
                "Bonus Tracker",
                "Emergency Fund",
                "5-Year Projection",
            ],
        )

    def test_expense_tables_have_unique_names(self) -> None:
        workbook = self.generate()
        names = [
            table.name
            for worksheet in workbook.worksheets
            for table in worksheet.tables.values()
        ]
        self.assertEqual(
            names,
            ["Partner1ExpenseTable", "Partner2ExpenseTable", "SharedExpenseTable"],
        )
        self.assertEqual(len(names), len(set(names)))

    def test_income_proportional_split_and_allocation_check_are_formula_driven(self) -> None:
        setup = self.generate()["Household Setup"]
        self.assertEqual(setup["C14"].value, ContributionMethod.INCOME_PROPORTIONAL.value)
        self.assertIn("C9/(C9+D9)", setup["C17"].value)
        self.assertEqual(setup["C18"].value, "=1-C17")
        self.assertEqual(setup["C19"].value, "=C17+C18")

    def test_partner_contribution_budget_is_allocated_and_actual_is_editable(self) -> None:
        workbook = self.generate()
        partner = workbook["Partner 1 Budget"]
        budget_formula = partner.cell(row=PARTNER_ROWS.household_contribution, column=3).value
        self.assertIn(f"'Shared Household'!C{SHARED_ROWS.required_funding}", budget_formula)
        self.assertIn("'Household Setup'!$C17", budget_formula)
        self.assertIsNone(partner.cell(row=PARTNER_ROWS.household_contribution, column=4).value)

    def test_consolidated_expenses_exclude_internal_household_contributions(self) -> None:
        dashboard = self.generate()["Household Dashboard"]
        total_expense_formula = dashboard["C11"].value
        self.assertIn(f"'Partner 1 Budget'!C{PARTNER_ROWS.total_expenses}", total_expense_formula)
        self.assertIn(f"'Partner 2 Budget'!C{PARTNER_ROWS.total_expenses}", total_expense_formula)
        self.assertIn(f"'Shared Household'!C{SHARED_ROWS.total_expenses}", total_expense_formula)
        self.assertNotIn(str(PARTNER_ROWS.household_contribution), total_expense_formula)

    def test_household_surplus_and_annual_savings_rate_use_consolidated_rows(self) -> None:
        dashboard = self.generate()["Household Dashboard"]
        self.assertEqual(dashboard["C15"].value, "=C8-C11-C14")
        year_budget_col = 3 + len(self.config.months) * 3
        savings_rate_formula = dashboard.cell(row=16, column=year_budget_col).value
        self.assertIn(f"{dashboard.cell(row=8, column=year_budget_col).coordinate}", savings_rate_formula)
        self.assertIn(f"{dashboard.cell(row=14, column=year_budget_col).coordinate}", savings_rate_formula)

    def test_month_count_drives_emergency_fund_average(self) -> None:
        six_month_config = replace(self.config, months=("Jan", "Feb", "Mar", "Apr", "May", "Jun"))
        workbook = self.generate(six_month_config)
        self.assertTrue(workbook["Savings Goals"]["C7"].value.endswith("/6"))
        self.assertEqual(workbook["Trend Analysis"].max_row, 12)

    def test_every_cross_sheet_formula_targets_an_existing_sheet(self) -> None:
        workbook = self.generate()
        available_sheets = set(workbook.sheetnames)
        missing_targets = []
        for worksheet in workbook.worksheets:
            for row in worksheet.iter_rows():
                for cell in row:
                    if isinstance(cell.value, str) and cell.value.startswith("="):
                        for sheet_name in re.findall(r"'([^']+)'!", cell.value):
                            if sheet_name not in available_sheets:
                                missing_targets.append((worksheet.title, cell.coordinate, sheet_name))
        self.assertEqual(missing_targets, [])

    def test_workbook_requests_full_recalculation_on_open(self) -> None:
        workbook = self.generate()
        self.assertEqual(workbook.calculation.calcMode, "auto")
        self.assertTrue(workbook.calculation.fullCalcOnLoad)
        self.assertTrue(workbook.calculation.forceFullCalc)


if __name__ == "__main__":
    unittest.main()
